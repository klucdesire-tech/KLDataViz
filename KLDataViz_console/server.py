# -*- coding: utf-8 -*-
"""
KLDataViz — Console d'utilisation de l'agent (serveur local).
Backend : proxy orchestrateUR CVLynk + moteur d'analyse (pandas) + exports PDF/Excel/HTML.
Sert l'interface web sur http://localhost:8642
"""
import http.server, socketserver, json, os, io, threading, time, tempfile, urllib.request, urllib.error, base64, sys, subprocess, traceback
import pandas as pd
import numpy as np
import openpyxl
from openpyxl.styles import Font, PatternFill

# --- Agent 2 MODÉLISATION (nettoyage, typage, imputation, doublons) ---
try:
    from modeling import clean_dataframe, bivariate_stats, to_dates
    MODELING_OK = True
except Exception as _mod_exc:
    MODELING_OK = False
    MODELING_ERROR = str(_mod_exc)
    def clean_dataframe(df):
        raise RuntimeError(MODELING_ERROR)
    def bivariate_stats(df):
        return {}

# --- Agent DATA INGESTION (ingestion sans altération de la source) ---
try:
    from ingestion import ingest as ingest_agent, Registry as IngestionRegistry
    INGESTION_OK = True
    INGESTION_ERROR = None
except Exception as _ingest_exc:  # la console reste utilisable même si l'agent d'ingestion est indisponible
    INGESTION_OK = False
    INGESTION_ERROR = str(_ingest_exc)

BASE = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, BASE)

# --- Agent ARCHITECTE DE DONNÉES (compréhension + structure cible) ---
try:
    import architect, data_architect
    ARCHITECT_ERROR = None
except Exception as _arch_exc:
    architect = None
    ARCHITECT_ERROR = str(_arch_exc)

import hashlib, hmac, re, secrets
import tenancy as T
import sandbox, netguard
import studio
import sources
import shares
import dhis2
import dictionnaire
if T.CLOUD:
    import cloud
from urllib.parse import urlparse, parse_qs
PORT = int(os.environ.get("KLDAVIZ_PORT", "8642"))
ORCHESTRATOR = os.environ.get("CVLYNK_ORCHESTRATOR", "https://api.cvlynk.com")


def _load_agent_token():
    """Token d'agent CVLynk — jamais en dur dans le code publié.
    Ordre : variable d'environnement CVLYNK_AGENT_TOKEN, puis fichier local agent_token.txt
    (ignoré par Git), puis ~/.cvlynk-orchestrator.json."""
    tok = os.environ.get("CVLYNK_AGENT_TOKEN")
    if tok:
        return tok.strip()
    local = os.path.join(BASE, "agent_token.txt")
    if os.path.exists(local):
        with open(local, encoding="utf-8") as fh:
            return fh.read().strip()
    cfg = os.path.join(os.path.expanduser("~"), ".cvlynk-orchestrator.json")
    if os.path.exists(cfg):
        try:
            with open(cfg, encoding="utf-8") as fh:
                return (json.load(fh) or {}).get("access_token", "")
        except Exception:
            pass
    return ""


AGENT_TOKEN = _load_agent_token()

# Dernière analyse, fichier en cours et dernier contrat d'ingestion : par utilisateur, voir tenancy.state().

def orchestrator(path, method="GET", body=None):
    if T.CLOUD:  # l'agent CVLynk n'est utilisé que par la console locale
        return 503, {"error": "Orchestrateur désactivé dans la version en ligne."}
    req = urllib.request.Request(ORCHESTRATOR + path, method=method)
    req.add_header("Authorization", "Bearer " + AGENT_TOKEN)
    req.add_header("Accept", "application/json")
    data = None
    if body is not None:
        data = json.dumps(body).encode("utf-8")
        req.add_header("Content-Type", "application/json")
    try:
        with urllib.request.urlopen(req, data=data, timeout=15) as r:
            return r.status, json.loads(r.read().decode() or "{}")
    except urllib.error.HTTPError as e:
        try:
            return e.code, json.loads(e.read().decode() or "{}")
        except Exception:
            return e.code, {}
    except Exception as e:
        return 0, {"error": str(e)}

def _auto_charts(df):
    """Génère un graphique adapté à chaque variable + nuages de points multi-variables."""
    charts = []
    numeric_cols = df.select_dtypes(include=[np.number]).columns

    # Nuages de points pour les paires numériques les plus corrélées (multi-variable)
    if len(numeric_cols) >= 2:
        try:
            cm = df[numeric_cols].corr()
            pairs = []
            for i in range(len(cm.columns)):
                for j in range(i):
                    pairs.append((str(cm.columns[i]), str(cm.columns[j]), float(cm.iloc[i, j])))
            pairs = sorted(pairs, key=lambda x: -abs(x[2]))[:3]
            for a, b, r in pairs:
                if abs(r) >= 0.3:
                    sub = df[[a, b]].dropna()
                    if len(sub) > 500:
                        sub = sub.sample(500, random_state=0)
                    if len(sub) > 5:
                        charts.append({
                            "title": f"{a} vs {b} (r={r:.2f})", "type": "scatter",
                            "labels": [str(x) for x in sub[a].tolist()],
                            "values": [[float(y) for y in sub[b].tolist()]],
                            "seriesNames": [b],
                        })
        except Exception:
            pass

    for col in df.columns:
        name = str(col)
        s = df[col]

        # Détection de date
        is_date = pd.api.types.is_datetime64_any_dtype(s)
        if not is_date and not pd.api.types.is_numeric_dtype(s):
            sample = s.dropna().head(30)
            if len(sample) >= 3:
                try:
                    parsed = to_dates(sample)
                    if parsed.notna().mean() >= 0.9:
                        is_date = True
                except Exception:
                    pass

        if is_date:
            try:
                dt = to_dates(s)
                per = dt.dt.to_period("M").astype(str).value_counts().sort_index()
                charts.append({"title": name, "type": "line",
                               "labels": list(per.index), "values": [list(map(int, per.values))]})
            except Exception:
                pass
            continue

        if col in numeric_cols:
            vals = s.dropna()
            nunique = int(vals.nunique())
            if nunique <= 15:
                vc = vals.value_counts().sort_index()
                charts.append({"title": name, "type": "bar",
                               "labels": [str(k) for k in vc.index], "values": [list(map(int, vc.values))]})
            else:
                try:
                    counts, edges = np.histogram(vals, bins=10)
                    labels = [f"{edges[i]:.3g}-{edges[i+1]:.3g}" for i in range(len(counts))]
                    charts.append({"title": name, "type": "histogram",
                                   "labels": labels, "values": [list(map(int, counts))]})
                except Exception:
                    pass
            continue

        vc = s.value_counts(dropna=True)
        if len(vc) == 0:
            continue
        if len(vc) > 20 and vc.iloc[0] <= 1:  # identifiant (une valeur par ligne) : un classement n'a pas de sens
            continue
        top = vc.head(12)
        charts.append({"title": name, "type": "bar",
                       "labels": [str(k)[:40] for k in top.index], "values": [list(map(int, top.values))]})

    return charts


def normalize_source_url(url):
    """Transforme un lien de partage (Google Sheets, Google Drive) en lien de téléchargement direct."""
    import re
    from urllib.parse import urlparse, parse_qs
    u = urlparse(url)
    if "docs.google.com" in u.netloc and "/spreadsheets/d/" in u.path:
        m = re.search(r"/spreadsheets/d/([^/]+)", u.path)
        if m and m.group(1) != "e":  # /d/e/... = lien « publié sur le web », déjà téléchargeable
            gid = (parse_qs(u.query).get("gid") or parse_qs(u.fragment).get("gid") or [None])[0]
            return (f"https://docs.google.com/spreadsheets/d/{m.group(1)}/export?format=csv"
                    + (f"&gid={gid}" if gid else ""))
    if "drive.google.com" in u.netloc:
        m = re.search(r"/file/d/([^/]+)", u.path) or re.search(r"[?&]id=([^&]+)", url)
        if m:
            return f"https://drive.google.com/uc?export=download&id={m.group(1)}"
    return url


def source_label(file_path):
    """Nom lisible et sans caractère interdit pour une source (fichier ou URL)."""
    import re
    from urllib.parse import urlparse
    if file_path.startswith(("http://", "https://")):
        u = urlparse(file_path)
        if "docs.google.com" in u.netloc and "/spreadsheets/" in u.path:
            m = re.search(r"/d/([^/]+)", u.path)
            name = "google_sheet_" + (m.group(1)[:8] if m else "")
        else:
            name = os.path.splitext(os.path.basename(u.path))[0] or u.netloc
    else:
        name = os.path.splitext(os.path.basename(file_path))[0]
    return re.sub(r'[\\/:*?"<>|]+', "_", name).strip("._ ") or "source"


def _csv_options(file_path):
    """Séparateur et encodage d'un CSV ; options par défaut de pandas si la détection échoue."""
    try:
        enc, sep = data_architect._sniff_text(data_architect.Path(file_path))
        return {"sep": sep, "encoding": enc}
    except Exception:
        return {}


_header_row = data_architect.header_row  # partagé avec l'architecte et le skill


def _tidy(df):
    """Noms de colonnes lisibles (en-têtes vides ou sur plusieurs lignes) et lignes vides retirées."""
    names, seen = [], {}
    for j, c in enumerate(df.columns):
        nm = " ".join(str(c).split())
        if not nm or nm.startswith("Unnamed:"):
            nm = f"colonne_{j + 1}"
        if nm in seen:
            seen[nm] += 1
            nm = f"{nm}_{seen[nm]}"
        else:
            seen[nm] = 1
        names.append(nm)
    df.columns = names
    return df.dropna(how="all").reset_index(drop=True)


AGREGE_SUFFIX = ".agrege"  # fichier témoin : « ces valeurs sont déjà agrégées » (ex. copie d'une source DHIS2)


def mark_agrege(path, on=True):
    """Pose (ou retire) le témoin à côté d'un fichier dérivé, pour qu'AGT-11 le sache après un aller-retour CSV."""
    side = path + AGREGE_SUFFIX
    if on:
        open(side, "w").close()
    elif os.path.exists(side):
        os.remove(side)


def _read_df(file_path, sheet=None):
    """Comme _read_df_raw, et réattache le drapeau « agrégé » d'un fichier dérivé (témoin .agrege)."""
    df, err = _read_df_raw(file_path, sheet)
    if df is not None and not file_path.startswith(("http://", "https://", "dhis2://")) \
            and os.path.exists(file_path + AGREGE_SUFFIX):
        df.attrs["agrege"] = True
    return df, err


def _read_df_raw(file_path, sheet=None):
    """Lit un fichier local ou une URL en DataFrame (csv/xlsx/xls/xlsm/json/parquet).

    Classeur Excel : la feuille `sheet`, sinon celle qui porte les données (pas forcément la
    première : page de garde, sommaire…). df.attrs["feuille"] / ["feuilles"] disent laquelle.
    En ligne : seuls les fichiers de l'organisation et les liens https publics sont acceptés."""
    if file_path.startswith("dhis2://"):
        return dhis2.read_dhis2(file_path, allow_private=not T.CLOUD)
    tmp_to_delete = None
    is_url = file_path.startswith(("http://", "https://"))
    if T.CLOUD and not is_url and not T.current().owns(file_path):
        return None, "Fichier introuvable."
    if is_url:
        url = normalize_source_url(file_path)
        try:
            data, ctype = netguard.fetch(url, allow_private=not T.CLOUD, https_only=T.CLOUD)
        except netguard.BlockedURL as e:
            return None, str(e)
        except Exception as e:
            if "google.com" in url and getattr(e, "code", None) in (401, 403, 404):
                return None, ("Google refuse l'accès à cette feuille. Vérifiez le lien, puis dans Google Sheets "
                              "cliquez sur Partager → « Tous les utilisateurs disposant du lien » (Lecteur).")
            return None, f"Téléchargement impossible : {e}"
        if "text/html" in ctype or data[:300].lstrip().lower().startswith((b"<!doctype html", b"<html")):
            if "google.com" in url:
                return None, ("Google renvoie une page de connexion au lieu des données. Dans Google Sheets, "
                              "cliquez sur Partager puis « Tous les utilisateurs disposant du lien » (Lecteur), "
                              "et réessayez avec le même lien.")
            return None, ("Ce lien renvoie une page web, pas un fichier de données. Utilisez un lien direct "
                          "vers un fichier CSV, Excel ou JSON.")
        ext = os.path.splitext(url.split("?")[0])[1].lower()
        if ext not in (".csv", ".xlsx", ".xls", ".xlsm", ".json", ".parquet"):
            ext = (".xlsx" if "spreadsheetml" in ctype or data[:2] == b"PK"
                   else ".json" if "json" in ctype or data[:1] in (b"{", b"[") else ".csv")
        tmp = tempfile.NamedTemporaryFile(suffix=ext, delete=False)
        tmp.write(data); tmp.close()
        file_path = tmp_to_delete = tmp.name
    if not os.path.exists(file_path):
        return None, f"Fichier introuvable : {file_path}"
    ext = os.path.splitext(file_path)[1].lower()
    try:
        if ext == ".csv":
            # Séparateur et encodage détectés : les exports français et DHIS2 utilisent « ; » (et parfois cp1252)
            opts = _csv_options(file_path)
            h = _header_row(pd.read_csv(file_path, header=None, nrows=20, dtype=str, on_bad_lines="skip", **opts))
            df = _tidy(pd.read_csv(file_path, skiprows=h, **opts)) if h else pd.read_csv(file_path, **opts)
        elif ext in (".xlsx", ".xls", ".xlsm"):
            raw, lue = data_architect.read_sheet(file_path, sheet)
            df = _tidy(raw)
            try:
                meta = {f["nom"]: f for f in data_architect.excel_sheets(file_path)}
            except Exception:
                meta = {}
            if meta.get(lue, {}).get("rapport"):
                df = data_architect.tidy_report(df)
            names = pd.ExcelFile(file_path).sheet_names
            if len(names) > 1:
                df.attrs.update(feuille=lue if isinstance(lue, str) else names[lue], feuilles=names)
        elif ext == ".json":
            df = pd.read_json(file_path)
        elif ext == ".parquet":
            df = pd.read_parquet(file_path)
        else:
            return None, f"Format non supporté : {ext}"
    except Exception as e:
        return None, f"Lecture impossible : {e}"
    finally:
        if tmp_to_delete:
            try:
                os.remove(tmp_to_delete)
            except OSError:
                pass
    return df, None


_PERIOD_HINT = None


def _libelles():
    d = os.path.join(os.path.dirname(os.path.dirname(os.path.abspath(__file__))), "HERMES_DA")
    if d not in sys.path:
        sys.path.insert(0, d)
    from hermes_da import libelles
    return libelles


def _pretty(name):
    """Libellé lisible d'une colonne (accents, sigles, unité) : « delai_initiation_jours » → « Délai initiation (jours) »."""
    try:
        return _libelles().joli(name)
    except Exception:
        s = str(name).replace("_", " ").strip()
        return s[:1].upper() + s[1:]


def _pretty_min(name):
    try:
        return _libelles().joli_min(name)
    except Exception:
        return str(name).replace("_", " ")


REPORT_COLS = {"indicateur", "niveau", "sexe", "tranche_age", "site", "valeur"}


def _is_report(df):
    return REPORT_COLS <= set(map(str, df.columns))


def _short(label, n=60):
    label = " ".join(str(label).split())
    return label if len(label) <= n else label[:n - 1] + "…"


def _report_charts(df):
    """Rapport remis à plat : chaque graphique reste à un seul niveau (total, sexe ou âge), jamais additionnés."""
    charts = []
    rnd = lambda xs: [round(float(x), 4) for x in xs]
    tot = df[df["niveau"] == "total"].groupby("indicateur", sort=False)["valeur"].sum()
    if len(tot) > 1:
        charts.append({"kind": "ranking", "type": "bar", "title": "Total par indicateur",
                       "subtitle": "Somme sur tous les sites", "labels": [_short(i) for i in tot.index],
                       "values": [rnd(tot.values)], "seriesNames": ["Total"], "horizontal": True})
    detail = [i for i in tot.sort_values(ascending=False).index
              if (df[(df["indicateur"] == i) & (df["niveau"] != "total")]["valeur"].sum() > 0)]
    for ind in detail[:2]:
        g = df[df["indicateur"] == ind]
        sx = g[g["niveau"] == "sexe"].groupby("sexe")["valeur"].sum()
        if len(sx) >= 2:
            charts.append({"kind": "report", "type": "bar", "title": f"{_short(ind, 45)} — par sexe",
                           "subtitle": f"Somme sur {g['site'].nunique()} site(s)", "labels": [str(x) for x in sx.index],
                           "values": [rnd(sx.values)], "seriesNames": ["Valeur"]})
        ag = g[g["niveau"] == "age"].pivot_table(index="tranche_age", columns="sexe", values="valeur", aggfunc="sum").fillna(0)
        if len(ag) >= 2:
            ag = ag.loc[sorted(ag.index, key=data_architect.age_order)]
            charts.append({"kind": "trend", "type": "bar", "title": f"{_short(ind, 45)} — par âge et sexe",
                           "subtitle": "Somme sur tous les sites", "labels": [str(x) for x in ag.index],
                           "values": [rnd(ag[c].values) for c in ag.columns], "seriesNames": [str(c) for c in ag.columns]})
    if len(tot):
        ind = tot.idxmax()
        st = df[(df["indicateur"] == ind) & (df["niveau"] == "total")].groupby("site")["valeur"].sum().sort_values(ascending=False)
        if len(st) >= 2:
            top = st.head(15)
            charts.append({"kind": "ranking", "type": "bar", "title": f"{_short(ind, 45)} — par site",
                           "subtitle": f"{len(top)} premiers sites sur {len(st)} · ligne : moyenne",
                           "labels": [str(x) for x in top.index], "values": [rnd(top.values)], "seriesNames": ["Valeur"],
                           "ref": {"value": round(float(st.mean()), 4), "label": "Moyenne"}, "horizontal": True})
    return charts


def report_checks(df):
    """Contrôle qualité d'un rapport : le total d'un indicateur doit égaler la somme par sexe et par âge,
    site par site. Renvoie les écarts trouvés."""
    out = []
    for (ind, site), g in df.groupby(["indicateur", "site"], sort=False):
        tot = g[g["niveau"] == "total"]["valeur"].sum()
        for niveau, nom in (("sexe", "par sexe"), ("age", "par âge")):
            part = g[g["niveau"] == niveau]
            if part.empty:
                continue
            if niveau == "age":  # chaque sexe doit égaler la somme de ses tranches d'âge
                sx = g[g["niveau"] == "sexe"].groupby("sexe")["valeur"].sum()
                ages = part.groupby("sexe")["valeur"].sum()
                for sexe in sx.index.intersection(ages.index):
                    if abs(sx[sexe] - ages[sexe]) > 1e-9:
                        out.append({"indicateur": ind, "site": site, "controle": f"{sexe} = somme des âges",
                                    "attendu": float(sx[sexe]), "trouve": float(ages[sexe])})
                continue
            if abs(tot - part["valeur"].sum()) > 1e-9:
                out.append({"indicateur": ind, "site": site, "controle": f"Total = somme {nom}",
                            "attendu": float(tot), "trouve": float(part["valeur"].sum())})
    return out


TOTAL_LABEL = re.compile(r"^\s*(sous[-\s]?)?total(e|s|aux)?(\s+g[ée]n[ée]ral)?\s*$|^\s*(ensemble|grand total)\s*$", re.I)
_MOIS = {"janv": 1, "jan": 1, "fev": 2, "fevr": 2, "feb": 2, "mars": 3, "mar": 3, "avr": 4, "apr": 4, "mai": 5, "may": 5,
         "juin": 6, "jun": 6, "juil": 7, "jul": 7, "aout": 8, "aug": 8, "sept": 9, "sep": 9, "oct": 10, "nov": 11,
         "dec": 12}


def _fr_period(label):
    """« Avril 2025 », « avr.-25 », « T1 2025 », « 2025 T3 » → date du début de période (NaT sinon)."""
    s = str(label).strip().lower()
    s = (s.replace("é", "e").replace("è", "e").replace("û", "u").replace("ô", "o"))
    m = re.match(r"^(?:t|q|trim(?:estre)?\s*)([1-4])[\s\-_/]*(\d{4})$|^(\d{4})[\s\-_/]*(?:t|q)([1-4])$", s)
    if m:
        q, y = (m.group(1), m.group(2)) if m.group(1) else (m.group(4), m.group(3))
        return pd.Timestamp(int(y), 3 * int(q) - 2, 1)
    m = re.match(r"^([a-z]+)\.?[\s\-_/']*(\d{2}|\d{4})$", s)
    if m:
        word, y = m.group(1), int(m.group(2))
        mo = _MOIS.get(word) or _MOIS.get(word[:4]) or _MOIS.get(word[:3])
        if mo:
            return pd.Timestamp(y + 2000 if y < 100 else y, mo, 1)
    return pd.NaT


def _cumul_rows(df):
    """Lignes de cumul (« Total »…) mêlées aux lignes de détail : nombre et libellés, pour avertir l'utilisateur."""
    txt = [c for c in df.columns if not pd.api.types.is_numeric_dtype(df[c])]
    if not txt:
        return {"lignes": 0, "libelles": []}
    hits = df[txt].astype(str).apply(lambda s: s.str.match(TOTAL_LABEL))
    mask = hits.any(axis=1)
    if not mask.any() or (~mask).sum() < 2:
        return {"lignes": 0, "libelles": []}
    labels = sorted({str(v).strip() for c in txt for v in df.loc[hits[c], c]})
    return {"lignes": int(mask.sum()), "libelles": labels[:5]}


def _infer_kind(c):
    """Graphiques « par variable » (_auto_charts) : type de lecture déduit du type de graphique."""
    if c.get("kind"):
        return c
    typ, title = c.get("type"), str(c.get("title") or "")
    if typ == "scatter":
        m = re.match(r"^(.*?) vs (.*?) \(r=(-?[\d.]+)\)$", title)
        return dict(c, kind="relation", subtitle=c.get("subtitle") or (f"r = {m.group(3)}" if m else ""),
                    xTitle=m.group(1) if m else "", yTitle=m.group(2) if m else "") if m else c
    if typ == "line":
        return dict(c, kind="trend", seriesNames=c.get("seriesNames") or [f"Le nombre d'enregistrements ({title})"])
    if typ == "histogram":
        return dict(c, kind="distribution")
    if typ in ("pie", "doughnut"):
        return dict(c, kind="share")
    if typ == "bar":
        labels = [str(x) for x in c.get("labels") or []]
        numeric = labels and all(re.fullmatch(r"-?[\d.,]+", x) for x in labels)
        return dict(c, kind="distribution" if numeric else "ranking")
    return c


def _with_rule_readings(charts):
    """Lecture chiffrée calculée par règles pour chaque graphique automatique (la même que dans le PDF) :
    affichée tout de suite, avant que Hermes ne rédige la sienne."""
    import html as _html
    try:
        import pdf_report
    except Exception:
        return charts
    out = []
    for c in charts:
        try:
            txt = pdf_report.interpret_chart(c)
        except Exception:
            txt = ""
        if txt:
            txt = _html.unescape(re.sub(r"</?b>", "**", txt))
            c = dict(c, lecture_auto=txt)
        out.append(c)
    return out


def _smart_charts(df, max_charts=6):
    """Graphiques « décisionnels » choisis automatiquement : évolution dans le temps, classement par
    groupe (avec ligne de moyenne), répartition, distribution et relation. Chaque graphique porte
    son agrégation (somme ou moyenne) pour que l'interface puisse l'expliquer."""
    import re
    # Les lignes de cumul (« Total », « Ensemble ») ne sont pas des catégories : elles doubleraient les sommes.
    txt = [c for c in df.columns if not pd.api.types.is_numeric_dtype(df[c])]
    if txt:
        cumul = df[txt].astype(str).apply(lambda s: s.str.match(TOTAL_LABEL)).any(axis=1)
        if cumul.any() and (~cumul).sum() >= 2:
            df = df[~cumul]
    # Réservations annulées, opérations rejetées : hors des montants et des volumes (sinon chiffre d'affaires gonflé)
    excl = ""
    for c in df.columns:
        if re.search(r"annul|cancel|rejet|refus", str(c), re.I):
            low = df[c].dropna().astype(str).str.strip().str.lower()
            yes = {"oui", "yes", "true", "vrai", "1", "y"}
            if low.isin(yes).any() and (~low.isin(yes)).any():
                df = df[~df[c].astype(str).str.strip().str.lower().isin(yes)]
                excl = f"Hors « {_pretty_min(c)} » · "
                break
    n = len(df)
    if n == 0:
        return []
    id_hint = re.compile(r"(^|[_\s])(id|code|key|ref|numero|num|uid)($|[_\s])", re.I)
    rate_hint = re.compile(r"(pct|percent|pourcent|taux|rate|ratio|couverture|%|moyenne|avg|mean|score)", re.I)
    period_hint = re.compile(r"(date|mois|month|semaine|week|annee|année|year|periode|période|period|trimestre|quarter|jour|day)", re.I)

    numeric_all = list(df.select_dtypes(include=[np.number]).columns)

    def is_year(s):
        v = s.dropna()
        return len(v) > 0 and v.between(1900, 2100).all() and (v % 1 == 0).all() and v.nunique() <= 60

    num = [c for c in numeric_all if not id_hint.search(str(c)) and df[c].nunique() > 1 and not is_year(df[c])]
    num.sort(key=lambda c: (-round(df[c].notna().mean(), 1), -float(df[c].abs().sum())))

    # Grandeurs « par unité » ou individuelles : un prix, un âge, une note, un délai, une charge virale ne
    # s'additionnent JAMAIS (la somme des prix unitaires ou des âges n'a pas de sens).
    intensive = re.compile(r"(prix|price|tarif|unitaire|_par_|_kg$|(^|_)kg($|_)|note|score|(^|_)age($|_)|âge|poids|"
                           r"weight|taille|height|imc|bmi|cd4|copies|charge_?virale|(^|_)cv($|_)|temperature|humidite|"
                           r"anciennete|delai|duree|latitude|longitude|densite|rendement)", re.I)

    def agg_of(c):
        v = df[c].dropna()
        frac = len(v) > 0 and v.between(0, 1).all() and (v % 1 != 0).any()
        if rate_hint.search(str(c)) or frac or intensive.search(str(c)):
            med = float(v.median()) if len(v) else 0.0
            # valeurs très dispersées (délais, charges virales) : la médiane résume mieux que la moyenne
            return "median" if len(v) and v.min() >= 0 and med > 0 and float(v.mean()) > 1.5 * med else "mean"
        return "sum"

    def lab(a):
        return {"sum": "Somme", "mean": "Moyenne", "median": "Médiane", "count": "Nombre de lignes"}[a]

    units = re.compile(r"(fcfa|xof|eur|usd|dollars?|ha|hectares?|kg|tonnes?|(?<=_)t$|jours?|nuits?|pct|pourcent|%|"
                       r"patients?|personnes?|habitants?|lignes?)", re.I)

    def unit_tokens(c):
        return {u.lower().rstrip("s") for u in units.findall(str(c))}

    # Volumes d'activité d'abord (consultations, ventes), puis les stocks (population desservie : un contexte, pas
    # l'activité suivie), puis les grandeurs individuelles — jamais choisis pour leur seule ampleur
    stock_words = re.compile(r"population|habitants|superficie|cheptel|capacite|solde|encours|file_?active|tx_?curr", re.I)
    num.sort(key=lambda c: 2 if agg_of(c) != "sum" else 1 if stock_words.search(str(c)) else 0)

    def unit_of(c):
        v = df[c].dropna()
        return "%" if re.search(r"(pct|percent|pourcent|%|taux|couverture)", str(c), re.I) and len(v) and v.between(0, 100).all() else ""

    # ---- colonne de période (date réelle ou libellé de période)
    # Une date « attribut » (naissance, embauche, création) décrit une personne ou un objet, pas l'activité suivie :
    # « l'évolution du salaire par date d'embauche » n'a pas de sens.
    attribut = re.compile(r"naissance|birth|embauche|hire|creation|cr[ée]ation|ouverture|expiration|p[ée]remption", re.I)
    period, period_dt = None, None
    for c in df.columns:
        if c in numeric_all or attribut.search(str(c)):
            continue
        s = df[c]
        if pd.api.types.is_datetime64_any_dtype(s):
            period, period_dt = c, s
            break
        sample = s.dropna().astype(str).head(40)
        if len(sample) >= 3:
            try:
                parsed = to_dates(sample)
                if parsed.notna().mean() >= 0.9:
                    period, period_dt = c, to_dates(s)
                    break
            except Exception:
                pass
            # Libellés en toutes lettres (« Avril 2025 », « T1 2025 ») : triés dans l'ordre du temps, pas de l'alphabet.
            fr = sample.map(_fr_period)
            if fr.notna().mean() >= 0.9:
                period, period_dt = c, pd.to_datetime(s.map(_fr_period))
                break
    if period is None:
        for c in df.columns:
            if c not in numeric_all and period_hint.search(str(c)) and 3 <= df[c].nunique() <= 60:
                period = c
                break

    group_hint = re.compile(r"(r[ée]gion|district|site|type|cat[ée]g|groupe|group|sexe|sex|genre|province|d[ée]partement|"
                            r"zone|structure|[ée]tablissement|facility|ville|commune|pays|country|service|produit|product|"
                            r"segment|canal|ip_?s?name|agence|magasin|client)", re.I)
    bool_vocab = {"true", "false", "vrai", "faux", "oui", "non", "yes", "no", "0", "1", "y", "n"}

    def is_bool(c):
        return set(df[c].dropna().astype(str).str.strip().str.lower().unique()) <= bool_vocab

    cats = [c for c in df.columns if c not in numeric_all and c != period and not id_hint.search(str(c))
            and 2 <= df[c].nunique() <= 40 and df[c].notna().mean() >= 0.6 and not is_bool(c)]
    # Noms « métier » d'abord, puis les regroupements de 3 à 20 valeurs, puis les colonnes « une ligne = une entité ».
    cats.sort(key=lambda c: (0 if group_hint.search(str(c)) else 1, 0 if 3 <= df[c].nunique() <= 20 else 1,
                             df[c].nunique() >= 0.9 * n, df[c].nunique()))

    charts = []

    def series(grouped):
        return [None if pd.isna(v) else round(float(v), 4) for v in grouped]

    # 1. Évolution
    if period is not None and num:
        # Mesure individuelle (CV, CD4, âge…) : sa médiane mensuelle sur ~60 lignes n'est que du bruit ;
        # l'évolution utile est le nombre d'enregistrements (patients dépistés, ventes…) par période.
        count_mode = agg_of(num[0]) != "sum"
        same = [] if count_mode else [c for c in num[1:] if agg_of(c) == agg_of(num[0]) and unit_tokens(c) == unit_tokens(num[0])]
        metrics = [] if count_mode else [num[0]] + same[:1]
        if period_dt is not None:
            span = (period_dt.max() - period_dt.min()).days if period_dt.notna().any() else 0
            freq = "D" if span <= 62 else "W" if span <= 200 else "M"
            # Le pas des données prime : des mois (« 2026-01 ») ne doivent pas être éclatés en semaines.
            u = period_dt.dropna().drop_duplicates().sort_values()
            gap = float(u.diff().dt.days.median()) if len(u) > 1 else 0
            if gap >= 85:
                freq = "Q"
            elif gap >= 28:
                freq = "M"
            elif gap >= 7 and freq == "D":
                freq = "W"
            key = period_dt.dt.to_period(freq)
        else:
            key = df[period].astype(str)
        g = df.assign(_k=key).dropna(subset=["_k"]).groupby("_k")
        vals, names, aggs = [], [], []
        if count_mode:
            vals.append(series(g.size()))
            names.append("Nombre d'enregistrements")
            aggs.append("count")
        for m in metrics:
            a = agg_of(m)
            vals.append(series(getattr(g[m], a)()))
            names.append(_pretty(m))
            aggs.append(a)
        labels = [str(k) for k in g.size().index]
        partial = ""
        if period_dt is not None and len(labels) >= 4:
            # Dernière période commencée mais pas finie (ex. données jusqu'au 13 du mois) et nettement moins remplie :
            # la garder ferait « reculer » toutes les courbes à tort.
            sizes = g.size()
            last_k = sizes.index[-1]
            try:  # dates au jour près (pas des mois « 2025-10-01 ») et dernière période non terminée
                unfinished = period_dt.dt.day.nunique() > 1 and period_dt.max() < last_k.end_time.normalize()
            except Exception:
                unfinished = False
            if unfinished and sizes.iloc[-1] < 0.75 * float(sizes.iloc[:-1].median()):
                labels = labels[:-1]
                vals = [v[:-1] for v in vals]
                partial = f" · {last_k} incomplet, écarté"
        if len(labels) >= 3:
            mx = [max([abs(x) for x in v if x is not None] or [0]) for v in vals]
            two_axes = len(vals) == 2 and min(mx) > 0 and max(mx) / min(mx) > 4
            charts.append({"kind": "trend", "type": "line",
                           "title": ("Évolution du nombre d'enregistrements" if count_mode else "Évolution de " + " et ".join(_pretty_min(m) for m in metrics)),
                           "subtitle": lab(aggs[0]) + " par " + _pretty_min(period) + partial, "agg": aggs[0],
                           "labels": labels, "values": vals, "seriesNames": names, "axis2": two_axes,
                           "unit": unit_of(metrics[0]) if metrics else ""})

    # Indicateurs de stock (file active, TX_CURR…) : jamais additionnés d'une période à l'autre ; les classements
    # et répartitions portent alors sur la dernière période.
    stock_hint = re.compile(r"(file_?active|tx_?curr|actifs?|sous_?(arv|traitement)|stock|effectif_?present|patients?_?suivis?|"
                            r"population|habitants|superficie|cheptel|parc_|capacite|solde|encours|charge_?virale|"
                            r"cv_?supprim|tx_?pvls|actuellement)", re.I)
    flux_hint = re.compile(r"nouve|initi|d[ée]c[èe]s|d[ée]pist", re.I)  # « nouvellement initiés sous TARV » : un flux
    last, last_txt = None, ""
    if period_dt is not None and period_dt.notna().any():
        last = df[period_dt == period_dt.max()]
        last_txt = "Dernière période (" + str(period_dt.max().to_period("M")) + ") · "
    else:
        year_col = next((c for c in numeric_all if is_year(df[c]) and df[c].nunique() >= 2), None)
        if year_col is not None:  # données annuelles (colonne « annee ») : le stock se lit sur la dernière année
            ymax = int(df[year_col].max())
            last, last_txt = df[df[year_col] == ymax], f"Année {ymax} · "

    def frame_for(m):
        stock = stock_hint.search(str(m)) and not flux_hint.search(str(m))
        return (last, last_txt) if last is not None and len(last) and stock else (df, "")

    # 2. Classements par groupe
    pairs = []
    if cats and num:
        pairs.append((cats[0], num[0]))
        if len(cats) > 1:
            pairs.append((cats[1], num[0]))
        if len(num) > 1:
            pairs.append((cats[0], num[1]))
    kept = 0
    for cat, m in pairs[:3]:
        if kept >= 2:
            break
        a = agg_of(m)
        src, pre = frame_for(m)
        g = getattr(src.groupby(cat)[m], a)().dropna().sort_values(ascending=False)
        if len(g) < 2 or (g != 0).sum() < 2:  # une seule catégorie non nulle : rien à comparer
            continue
        if a != "sum" and float(g.max() - g.min()) <= 0.02 * abs(float(g.mean())):
            continue  # moyennes quasi identiques (87,6 vs 87,5) : un graphique sans information
        kept += 1
        top = g.head(12)
        charts.append({"kind": "ranking", "type": "bar", "title": f"{_pretty(m)} par {_pretty_min(cat)}",
                       "subtitle": pre + lab(a) + (f" · {len(top)} premiers sur {len(g)}" if len(g) > 12 else "")
                       + " · ligne : moyenne", "agg": a, "labels": [str(x) for x in top.index], "values": [series(top.values)],
                       "seriesNames": [_pretty(m)], "ref": {"value": round(float(g.mean()), 4), "label": "Moyenne"},
                       "horizontal": True, "unit": unit_of(m)})

    # 2 bis. Taux d'une colonne Oui/Non par groupe (admission, initiation le jour même…)
    yes_vals = {"oui", "yes", "true", "vrai", "1", "y"}
    bools = [c for c in df.columns if c not in numeric_all and df[c].nunique() == 2 and is_bool(c)
             and df[c].astype(str).str.strip().str.lower().isin(yes_vals).any()]
    if bools and cats:
        b, cat = bools[0], cats[0]
        flag = df[b].astype(str).str.strip().str.lower().isin(yes_vals).where(df[b].notna())
        g = (flag.groupby(df[cat]).mean() * 100).dropna().sort_values(ascending=False)
        if len(g) >= 2 and float(g.max() - g.min()) >= 2:  # moins de 2 points d'écart : rien à montrer
            top = g.head(12)
            charts.append({"kind": "ranking", "type": "bar", "title": f"Taux « {_pretty_min(b)} » par {_pretty_min(cat)}",
                           "subtitle": "Part de « Oui » · ligne : ensemble", "agg": "mean", "unit": "%",
                           "labels": [str(x) for x in top.index], "values": [series(top.values)],
                           "seriesNames": [f"Taux « {_pretty_min(b)} »"], "horizontal": True,
                           "ref": {"value": round(float(flag.mean() * 100), 4), "label": "Ensemble"}})

    # 3. Répartition
    small = [c for c in cats if 2 <= df[c].nunique() <= 6]
    if small:
        c = small[0]
        src, pre = frame_for(num[0]) if num else (df, "")
        if num and agg_of(num[0]) == "sum":
            g = src.groupby(c)[num[0]].sum().sort_values(ascending=False)
            title = f"Répartition de {_pretty_min(num[0])} par {_pretty_min(c)}"
        else:
            g = df[c].value_counts()
            title = f"Répartition des lignes par {_pretty_min(c)}"
        if len(g) >= 2 and float(g.sum()) > 0 and float(g.max()) / float(g.sum()) < 0.98:
            charts.append({"kind": "share", "type": "doughnut", "title": title, "subtitle": pre + "Part du total",
                           "labels": [str(x) for x in g.index], "values": [series(g.values)]})

    # 4. Distribution
    if num:
        m = num[0]
        v = df[m].dropna()
        heavy = len(v) and v.min() >= 0 and float(v.quantile(0.99)) > 50 * max(float(v.median()), 1.0)
        if v.nunique() >= 8 and heavy:
            # Valeurs très étalées (charge virale, montants) : des tranches égales mettraient tout dans la 1re barre
            top = int(np.ceil(np.log10(float(v.max()) + 1)))
            edges = [0, 1] + [10 ** k for k in range(1, top + 1)]
            counts = [int(((v >= edges[i]) & (v < edges[i + 1])).sum()) for i in range(len(edges) - 1)]
            counts[-1] += int((v >= edges[-1]).sum())

            def _p(x):
                return f"{int(x):,}".replace(",", " ")
            labels = ["0"] + [f"{_p(edges[i])} – {_p(edges[i + 1] - 1)}" for i in range(1, len(edges) - 1)]
            charts.append({"kind": "distribution", "type": "histogram", "title": f"Distribution de {_pretty_min(m)}",
                           "subtitle": "Nombre de lignes par ordre de grandeur (valeurs très étalées)",
                           "labels": labels, "values": [counts], "seriesNames": ["Lignes"]})
        elif v.nunique() >= 8:
            lo, hi = v.quantile(0.01), v.quantile(0.99)
            v2 = v[(v >= lo) & (v <= hi)]
            raw_step = (float(v2.max()) - float(v2.min())) / 10 or 1.0
            mag = 10 ** np.floor(np.log10(raw_step))
            step = next(m * mag for m in (1, 2, 2.5, 5, 10) if m * mag >= raw_step)  # pas « rond »
            start = np.floor(float(v2.min()) / step) * step
            edges = np.arange(start, float(v2.max()) + step, step)
            if len(edges) < 3:
                edges = np.array([start, start + step, start + 2 * step])
            counts, edges = np.histogram(v2, bins=edges)

            def _f(x):
                x = round(float(x), 6)
                return (f"{x:,.0f}" if x == int(x) else f"{x:,.2f}".rstrip("0")).replace(",", " ").replace(".", ",")
            charts.append({"kind": "distribution", "type": "histogram", "title": f"Distribution de {_pretty_min(m)}",
                           "subtitle": "Nombre de lignes par tranche" + (" · valeurs extrêmes (1 %) écartées" if len(v2) < len(v) else ""),
                           "labels": [f"{_f(edges[i])} – {_f(edges[i + 1])}" for i in range(len(counts))],
                           "values": [[int(x) for x in counts]], "seriesNames": ["Lignes"]})

    # 5. Relation (corrélation nette mais pas triviale)
    if len(num) >= 2:
        sub = df[num[:12]]
        cm = sub.corr()
        best = None
        for i in range(len(cm.columns)):
            for j in range(i):
                r = cm.iloc[i, j]
                ta = {t for t in re.split(r"[_\s]+", str(cm.columns[i]).lower()) if len(t) >= 4}
                tb = {t for t in re.split(r"[_\s]+", str(cm.columns[j]).lower()) if len(t) >= 4}
                derived = bool((ta & tb) - {"fcfa", "nombre", "total", "taux", "moyenne"})
                if pd.notna(r) and 0.5 <= abs(r) < 0.97 and not derived and (best is None or abs(r) > abs(best[2])):
                    best = (cm.columns[j], cm.columns[i], float(r))
        if best:
            a, b, r = best
            pts = df[[a, b]].dropna()
            if len(pts) > 400:
                pts = pts.sample(400, random_state=0)
            charts.append({"kind": "relation", "type": "scatter", "title": f"{_pretty(a)} et {_pretty_min(b)}",
                           "subtitle": f"Chaque point est une ligne · corrélation r = {r:.2f}".replace(".", ",", 1),
                           "labels": [float(x) for x in pts[a]], "values": [[float(y) for y in pts[b]]],
                           "seriesNames": [_pretty(b)], "xTitle": _pretty(a), "yTitle": _pretty(b), "trend": True})

    if excl:
        for c in charts:
            c["subtitle"] = excl + str(c.get("subtitle") or "")
    return charts[:max_charts]


_HEAVY = threading.BoundedSemaphore(int(os.environ.get("KLDV_MAX_HEAVY", "2")))
_heavy_depth = threading.local()
_HEAVY_WAIT = float(os.environ.get("KLDV_HEAVY_WAIT", "90"))  # secondes d'attente avant « serveur occupé »


def _lourd(fn):
    """Analyse, tableau de bord, calcul du chat : au plus KLDV_MAX_HEAVY à la fois (2 par défaut). Au-delà, on attend
    jusqu'à KLDV_HEAVY_WAIT s (90 par défaut) puis on répond « occupé » (HTTP 503) plutôt que de saturer la mémoire et de faire tomber le service."""
    import functools

    @functools.wraps(fn)
    def wrap(*a, **k):
        depth = getattr(_heavy_depth, "n", 0)
        if depth:  # déjà dans un traitement lourd (même fil) : pas de second jeton
            return fn(*a, **k)
        if not _HEAVY.acquire(timeout=_HEAVY_WAIT):
            raise T.Unavailable("Le serveur traite déjà d'autres analyses. Réessayez dans une minute.")
        _heavy_depth.n = 1
        try:
            return fn(*a, **k)
        finally:
            _heavy_depth.n = 0
            _HEAVY.release()
    return wrap


@_lourd
def analyze(file_path):
    df, err = _read_df(file_path)
    if err and re.search(r"Permission denied|Errno 13|being used by another process", str(err)):
        return {"error": "Fichier illisible : il est sans doute ouvert dans Excel. Fermez-le et réessayez."}
    if err:
        return {"error": err}
    if df is None:
        return {"error": "Fichier illisible (peut-être ouvert dans Excel ou protégé). Fermez-le et réessayez."}
    lu = dict(df.attrs)  # feuille lue dans un classeur
    libelles = {str(c): _pretty(c) for c in df.columns}

    size = os.path.getsize(file_path) if os.path.exists(file_path) else 0
    rows, cols = df.shape
    total_cells = rows * cols

    columns = []
    missing_total = 0
    for c in df.columns:
        na = int(df[c].isna().sum())
        missing_total += na
        columns.append({
            "name": str(c),
            "dtype": str(df[c].dtype),
            "missing": na,
            "missing_pct": round(na / rows * 100, 2) if rows else 0,
            "unique": int(df[c].nunique(dropna=True)),
        })

    numeric = df.select_dtypes(include=[np.number])
    num_stats = {}
    anomalies = []
    for c in numeric.columns:
        s = numeric[c].dropna()
        if len(s) == 0:
            continue
        q1, q3 = s.quantile(0.25), s.quantile(0.75)
        iqr = q3 - q1
        lo, hi = q1 - 1.5 * iqr, q3 + 1.5 * iqr
        outliers = s[(s < lo) | (s > hi)]
        skip = (re.search(r"(^|[_\s])(id|code|num|numero|matricule|annee|year|mois)($|[_\s])", str(c), re.I)
                or iqr == 0)  # identifiant, année, ou moitié des valeurs identiques (ex. naissances = 0) : pas d'IQR
        if not skip and s.min() >= 0 and float(s.quantile(0.99)) > 20 * max(float(s.median()), 1.0):
            # Données très asymétriques (montants, charges virales) : l'IQR brut marquerait des valeurs normales
            ls = np.log1p(s)
            lq1, lq3 = ls.quantile(0.25), ls.quantile(0.75)
            outliers = s[(ls < lq1 - 3 * (lq3 - lq1)) | (ls > lq3 + 3 * (lq3 - lq1))]
            lo, hi = float(np.expm1(lq1 - 3 * (lq3 - lq1))), float(np.expm1(lq3 + 3 * (lq3 - lq1)))
        if skip or len(outliers) > 0.05 * len(s):
            # Plus de 5 % des valeurs « hors plage » : ce sont des groupes différents (Suites vs Standard,
            # électroménager vs alimentation), pas des erreurs de saisie
            outliers = s.iloc[0:0]
        lo, hi = max(lo, float(s.min())), min(hi, float(s.max()))
        num_stats[str(c)] = {
            "mean": round(float(s.mean()), 4),
            "std": round(float(s.std()), 4),
            "min": round(float(s.min()), 4),
            "q25": round(float(q1), 4),
            "median": round(float(s.median()), 4),
            "q75": round(float(q3), 4),
            "max": round(float(s.max()), 4),
            "outliers": int(len(outliers)),
        }
        if len(outliers) > 0:
            anomalies.append({
                "column": str(c), "type": "valeur aberrante (IQR)",
                "count": int(len(outliers)),
                "detail": f"{len(outliers)} valeur(s) hors de la plage habituelle [{lo:,.2f} ; {hi:,.2f}]".replace(",", " "),
            })

    cat = df.select_dtypes(exclude=[np.number])
    cat_top = {}
    for c in cat.columns:
        vc = df[c].value_counts(dropna=True).head(5)
        cat_top[str(c)] = {str(k): int(v) for k, v in vc.items()}

    corr_pairs = []
    if numeric.shape[1] >= 2:
        cm = numeric.corr()
        for i in range(len(cm.columns)):
            for j in range(i):
                corr_pairs.append((str(cm.columns[i]), str(cm.columns[j]), float(cm.iloc[i, j])))
        corr_pairs = sorted(corr_pairs, key=lambda x: -abs(x[2]))[:10]

    completeness = round((1 - missing_total / total_cells) * 100, 2) if total_cells else 100

    return {
        "file": os.path.basename(file_path) if os.path.exists(file_path) else source_label(file_path),
        "libelles": libelles,
        "size_kb": round(size / 1024, 1),
        "rows": int(rows),
        "columns": int(cols),
        "completeness_pct": completeness,
        "missing_total": int(missing_total),
        "columns_detail": columns,
        "numeric_stats": num_stats,
        "categorical_top": cat_top,
        "correlations": [{"a": a, "b": b, "r": round(r, 3)} for a, b, r in corr_pairs],
        "anomalies": anomalies,
        "bivariate": bivariate_stats(df),
        "lignes_cumul": _cumul_rows(df),
        "charts": [] if _is_report(df) else _with_rule_readings([_infer_kind(c) for c in _auto_charts(df)]),
        "smart_charts": _with_rule_readings(_report_charts(df) if _is_report(df) else _smart_charts(df)),
        **({"controles": report_checks(df)[:200]} if _is_report(df) else {}),
        **({"feuille": {"lue": lu["feuille"], "autres": [n for n in lu["feuilles"] if n != lu["feuille"]]}}
           if lu.get("feuilles") else {}),
    }


def _cell_safe(v):
    """Convertit une cellule DataFrame en valeur JSON-safe."""
    if pd.isna(v):
        return None
    if isinstance(v, pd.Timestamp):
        return v.strftime("%Y-%m-%d")
    if isinstance(v, bool):
        return bool(v)
    try:
        if hasattr(v, "item"):
            return v.item()
    except Exception:
        pass
    return v


def clean(file_path, skip=None, hermes=None):
    """Agent 2 — modélisation : nettoie, met en forme et sauvegarde un fichier.
    `skip` : corrections refusées par l'utilisateur (voir modeling.clean_dataframe).
    `hermes` : corrections de Hermes acceptées (hermes_clean), appliquées AVANT les règles."""
    df, err = _read_df(file_path)
    if err:
        return {"error": err}
    agrege = bool(df.attrs.get("agrege"))
    h_details, h_applied = {}, []
    if hermes:
        import hermes_clean as HCL
        h_applied = [c for c in hermes if isinstance(c, dict)][:20]
        df, h_details, keep_text = HCL.apply(df, h_applied)
        skip = list(skip or []) + [f"type:{c}" for c in keep_text]  # codes gardés en texte
        # Valeurs que Hermes vient de vider (vides déguisés, valeurs impossibles) : ne pas les recompléter
        skip += [f"imput:{c.get('colonne')}" for c in h_applied
                 if c.get("type") == "plage" or (c.get("type") == "remplacer" and c.get("par") is None)]
    df_clean, rapport = clean_dataframe(df, skip=skip) if skip else clean_dataframe(df)
    if h_applied:
        rapport["hermes"] = [dict(c, detail=h_details.get(c.get("id"), {})) for c in h_applied]

    # Sauvegarde de la version nettoyée (c'est elle, et elle seule, qui sera ingérée)
    cleaned_dir = T.cleaned_dir()
    os.makedirs(cleaned_dir, exist_ok=True)
    stem = source_label(file_path)
    cleaned_path = os.path.join(cleaned_dir, f"{stem}_nettoye.csv")
    df_clean.to_csv(cleaned_path, index=False, encoding="utf-8-sig")
    mark_agrege(cleaned_path, agrege)  # le CSV perd df.attrs : le témoin garde l'information

    preview = df_clean.head(30)
    apercu = [[_cell_safe(x) for x in row]
              for row in preview.itertuples(index=False, name=None)]

    return {
        "file": source_label(file_path),
        "cleaned_path": cleaned_path,
        "rows": int(len(df_clean)),
        "cols": int(len(df_clean.columns)),
        "colonnes": [str(c) for c in df_clean.columns],
        "types": {str(c): str(df_clean[c].dtype) for c in df_clean.columns},
        "apercu": apercu,
        "rapport": rapport,
    }

def fetch_source():
    code, tasks = orchestrator("/api/v1/agents/tasks")
    if code != 200 or not isinstance(tasks, list) or not tasks:
        return {"source": None, "message": "Aucune tâche assignée pour le moment."}
    task = tasks[0]
    payload = task.get("payload") or {}
    source = (payload.get("source") or payload.get("file") or payload.get("link")
              or payload.get("url") or task.get("source") or task.get("file"))
    if source:
        return {"source": source, "task_id": task.get("id") or task.get("task_id")}
    return {"source": None, "message": "La tâche assignée ne contient pas de lien source."}

def _hermes_tasks():
    """Tâches de l'orchestrateur interne HERMES_DA (AGT-00) : jobs maîtres de planification."""
    try:
        from ingestion.store import default_dsn
        import psycopg
        conn = psycopg.connect(default_dsn(), autocommit=True, connect_timeout=5)
        try:
            with conn.cursor() as cur:
                cur.execute(
                    "SELECT job_id, objective, status, input_reference, created_at "
                    "FROM hermes.jobs WHERE task_type = 'plan' "
                    "ORDER BY created_at DESC LIMIT 50")
                rows = cur.fetchall()
        finally:
            conn.close()
    except Exception:
        return []
    tasks = []
    for job_id, objective, status, input_ref, created_at in rows:
        src = ""
        if isinstance(input_ref, dict):
            src = input_ref.get("source") or input_ref.get("path") or ""
        tasks.append({
            "type": (objective or "Analyse de données").strip(),
            "id": job_id,
            "task_id": job_id,
            "status": status or "DONE",
            "source": src,
            "created_at": str(created_at) if created_at else "",
        })
    return tasks

def dhis2_connect(p):
    """Lien API DHIS2 (copié depuis DHIS2) + identifiants → source « dhis2:// » prête à analyser.
    Le mot de passe est chiffré dans l'espace et n'est jamais renvoyé ni écrit dans les journaux."""
    link, user, pwd = str(p.get("link") or "").strip(), str(p.get("user") or "").strip(), str(p.get("password") or "")
    if T.CLOUD:
        cloud.require(T.current(), "owner", "analyst")
    if not link or not user or not pwd:
        raise ValueError("Indiquez le lien API DHIS2, l'utilisateur et le mot de passe.")
    if link.startswith("https://"):
        link = link[len("https://"):]
    elif link.startswith(dhis2.SCHEME):
        link = link[len(dhis2.SCHEME):]
    else:
        raise ValueError("Le lien DHIS2 doit commencer par https://")
    uri = dhis2.SCHEME + link
    instance, _ = dhis2.parse_uri(uri)
    df, err = dhis2.read_dhis2(uri, allow_private=not T.CLOUD, creds=(user, pwd))  # essai d'abord
    if err:
        return {"error": err}                       # mauvais mot de passe : rien n'est enregistré
    dhis2.save_credentials(instance, user, pwd)
    return {"uri": uri, "instance": instance, "name": f"DHIS2 {instance.split('/')[0]}", "rows": int(len(df)),
            "indicateurs": int(df["indicateur"].nunique()), "periodes": sorted(df["periode"].dropna().astype(str).unique())[-3:]}


@_lourd
def build_dashboard(df, spec=None):
    """AGT-11 sur un DataFrame. `spec` fournie (source connectée qu'on actualise) : mêmes indicateurs, sans IA."""
    hermes_da_dir = os.path.join(os.path.dirname(os.path.dirname(os.path.abspath(__file__))), "HERMES_DA")
    if hermes_da_dir not in sys.path:
        sys.path.insert(0, hermes_da_dir)
    from hermes_da.adapters.dashboard import generate_dashboard
    try:
        dico = dictionnaire.charger()  # dictionnaire d'indicateurs de l'espace ([] = normes intégrées)
    except Exception as exc:
        print(f"[dictionnaire] lecture impossible : {exc}", flush=True)
        dico = []
    if spec is not None:
        return generate_dashboard(df, "Tableau de bord KLDataViz", llm=False, spec=spec, dictionnaire=dico)
    cfg = _llm_config()  # en ligne : clé DEEPSEEK_API_KEY de l'environnement (pas de llm_config.json)
    model = os.environ.get("KLDV_DASHBOARD_MODEL") or (cfg or {}).get("dashboard_model") or "deepseek-flash"
    llm = (lambda m, n: _llm_call(m, n, model=model, effort="high", timeout=120, attempts=1)
           or _llm_call(m, n, model=model, effort="low", timeout=90, attempts=1)) if cfg else None
    return generate_dashboard(df, "Tableau de bord KLDataViz", llm=llm, dictionnaire=dico)


def do_dashboard(file_path, name=None):
    """Génère un tableau de bord interactif (AGT-11). Une source qu'on peut relire (lien Google Sheets / web,
    fichier local en mode local) devient une « source connectée » : actualisable sans réimporter."""
    try:
        df, err = _read_df(file_path)
        if err:
            return {"error": err}
        if df is None or len(df) == 0:
            return {"error": "Aucune donnée lisible dans cette source."}
        result = build_dashboard(df)
        out = {"html": result["html"], "ia": result["ia"], "kpis": result["kpis"]}
        try:
            out["source"] = sources.connect(file_path, name, result.get("spec"), df, result["html"])
        except Exception as exc:  # la connexion est un plus : le tableau de bord reste affiché
            print(f"[sources] connexion impossible : {exc}", flush=True)
        return out
    except Exception as e:
        return {"error": f"Génération du tableau de bord échouée : {e}"}

def export_excel(result):
    wb = openpyxl.Workbook()
    hdr = PatternFill("solid", fgColor="1286B4"); hdr_font = Font(bold=True, color="FFFFFF")
    def sheet(ws, title, headers, rows):
        ws.title = title
        for j, h in enumerate(headers, 1):
            c = ws.cell(1, j, h); c.fill = hdr; c.font = hdr_font
        for i, row in enumerate(rows, 2):
            for j, v in enumerate(row, 1):
                ws.cell(i, j, v)
    ws = wb.active
    sheet(ws, "Synthèse", ["Fichier", "Taille (Ko)", "Lignes", "Colonnes", "Complétude %", "Manquantes", "Anomalies"],
          [[result.get("file"), result.get("size_kb"), result.get("rows"), result.get("columns"),
            result.get("completeness_pct"), result.get("missing_total"), len(result.get("anomalies", []))]])
    sheet(wb.create_sheet(), "Colonnes", ["Variable", "Type", "Uniques", "Manquantes", "% manquant"],
          [[c["name"], c["dtype"], c["unique"], c["missing"], c["missing_pct"]] for c in result.get("columns_detail", [])])
    sheet(wb.create_sheet(), "Statistiques", ["Variable", "Moyenne", "Écart-type", "Min", "Médiane", "Max", "Hors-normes"],
          [[k, v["mean"], v["std"], v["min"], v["median"], v["max"], v["outliers"]] for k, v in result.get("numeric_stats", {}).items()])
    sheet(wb.create_sheet(), "Corrélations", ["Variable A", "Variable B", "r"],
          [[c["a"], c["b"], c["r"]] for c in result.get("correlations", [])])
    sheet(wb.create_sheet(), "Anomalies", ["Colonne", "Type", "Nombre", "Détail"],
          [[a["column"], a["type"], a["count"], a["detail"]] for a in result.get("anomalies", [])])
    fn = f"analyse_{time.strftime('%Y%m%d_%H%M%S')}.xlsx"
    wb.save(os.path.join(T.exports_dir(), fn))
    return fn

def export_html(result):
    def esc(s): return str(s).replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;")
    cols_html = "".join(f"<tr><td>{esc(c['name'])}</td><td>{esc(c['dtype'])}</td><td>{c['unique']}</td><td>{c['missing']}</td><td>{c['missing_pct']}%</td></tr>"
                        for c in result.get("columns_detail", []))
    stats_html = "".join(f"<tr><td>{esc(k)}</td><td>{v['mean']}</td><td>{v['std']}</td><td>{v['min']}</td><td>{v['median']}</td><td>{v['max']}</td><td>{v['outliers']}</td></tr>"
                         for k, v in result.get("numeric_stats", {}).items())
    corr_html = "".join(f"<tr><td>{esc(c['a'])}</td><td>{esc(c['b'])}</td><td>{c['r']}</td></tr>" for c in result.get("correlations", []))
    anom_html = "".join(f"<li><b>{esc(a['column'])}</b> — {esc(a['type'])} ({a['count']}) : {esc(a['detail'])}</li>" for a in result.get("anomalies", []))
    html = f"""<!DOCTYPE html><html lang="fr"><head><meta charset="utf-8"><title>Analyse — {esc(result.get('file'))}</title>
<style>body{{font-family:Segoe UI,Arial,sans-serif;max-width:900px;margin:32px auto;padding:0 20px;color:#1a1a1a;}}
h1{{color:#1286b4;}} table{{border-collapse:collapse;width:100%;font-size:13px;margin:10px 0;}}
th,td{{border:1px solid #cbd5e0;padding:6px 9px;text-align:left;}} th{{background:#1286b4;color:#fff;}}
.sum{{display:flex;gap:18px;flex-wrap:wrap;}} .kpi{{background:#f0f7fb;border:1px solid #cbe4f2;border-radius:8px;padding:12px 18px;}}
.kpi b{{font-size:22px;color:#1286b4;display:block;}} li{{margin:4px 0;}}</style></head><body>
<h1>Analyse — {esc(result.get('file'))}</h1>
<div class="sum"><div class="kpi"><b>{result.get('rows'):,}</b>lignes</div><div class="kpi"><b>{result.get('columns')}</b>colonnes</div>
<div class="kpi"><b>{result.get('completeness_pct')}%</b>complétude</div><div class="kpi"><b>{len(result.get('anomalies', []))}</b>anomalies</div></div>
<h2>Colonnes</h2><table><tr><th>Variable</th><th>Type</th><th>Uniques</th><th>Manquantes</th><th>%</th></tr>{cols_html}</table>
<h2>Statistiques numériques</h2><table><tr><th>Variable</th><th>Moyenne</th><th>Écart-type</th><th>Min</th><th>Médiane</th><th>Max</th><th>Hors-normes</th></tr>{stats_html}</table>
<h2>Corrélations</h2><table><tr><th>A</th><th>B</th><th>r</th></tr>{corr_html}</table>
<h2>Anomalies</h2><ul>{anom_html or '<li>Aucune anomalie détectée.</li>'}</ul>
<p style="color:#888;font-size:11px;">Généré par KLDataViz — agent Data Analytics &amp; BI</p></body></html>"""
    fn = f"analyse_{time.strftime('%Y%m%d_%H%M%S')}.html"
    with open(os.path.join(T.exports_dir(), fn), "w", encoding="utf-8") as f:
        f.write(html)
    return fn

def export_pdf(result):
    """Rapport détaillé : fiche du fichier, synthèse, qualité, graphiques interprétés, statistiques, relations."""
    import pdf_report
    fn = f"analyse_{time.strftime('%Y%m%d_%H%M%S')}.pdf"
    pdf_report.build(result, os.path.join(T.exports_dir(), fn))
    return fn

# ---------- Chat / interrogation des données en langage naturel ----------
def _chat_load_work(path):
    if not path:
        return None
    p = str(path)
    cands = []
    if os.path.isfile(p):
        cands = [p]
    elif os.path.isdir(p):
        cands = [os.path.join(p, f) for f in sorted(os.listdir(p)) if f.endswith((".parquet", ".csv"))][:1]
    else:
        for ext in (".parquet", ".csv"):
            if os.path.exists(p + ext):
                cands.append(p + ext)
    for cand in cands:
        try:
            if cand.endswith(".parquet"):
                return pd.read_parquet(cand)
            return pd.read_csv(cand)
        except Exception:
            continue
    return None

def chat_query(question):
    import re
    datasets = _registry_datasets()
    ql = (question or "").strip().lower()
    if not ql:
        return {"answer": "Posez une question sur les données (ex. « combien de lignes dans fact_district ? », « somme de value par district »).", "type": "text"}

    # --- liste des datasets
    if re.search(r"\b(liste|lister|quels|quelles|registre|datasets)\b", ql) and not any(
            (d.get("dataset_name") or "").lower() in ql for d in datasets):
        rows = [[d["dataset_id"], d["dataset_name"], d["row_count"], d["column_count"], d["source_format"]]
                for d in datasets]
        return {"answer": f"{len(datasets)} dataset(s) ingéré(s) :", "type": "table",
                "table": {"columns": ["dataset_id", "nom", "lignes", "colonnes", "format"], "rows": rows}}

    if not datasets:
        return {"answer": "Aucun dataset ingéré pour le moment. Ingérez d'abord un fichier (panneau Ingestion).", "type": "text"}

    # --- dataset cible : par nom, sinon par recouvrement de colonnes, sinon le plus récent
    reg = T.registry()  # joignable : il vient de fournir `datasets`
    dataset_cols = {}
    for d in datasets:
        try:
            dataset_cols[d["dataset_id"]] = [str(c.get("column_name") or "").lower()
                                             for c in reg.columns_of(d.get("dataset_version_id"))]
        except Exception:
            dataset_cols[d["dataset_id"]] = []
    target = None
    for d in datasets:
        nm = (d.get("dataset_name") or "").lower()
        if nm and nm in ql:
            target = d
            break
    if target is None:
        best = 0
        for d in datasets:
            score = sum(1 for c in dataset_cols.get(d["dataset_id"], []) if c and c in ql)
            if score > best:
                best, target = score, d
    if target is None:
        target = datasets[0]

    df = _chat_load_work(target.get("work_path"))
    if df is None:
        return {"answer": f"Impossible de charger les données de « {target.get('dataset_name')} ».", "type": "text"}

    cols = [str(c) for c in df.columns]
    colmap = {c.lower(): c for c in cols}
    mentioned = [colmap[k] for k in colmap if k and k in ql]
    numeric = [c for c in cols if str(df[c].dtype).startswith(("int", "float"))]

    # --- dimensions
    if re.search(r"combien de ligne|combien d.enregistrement|nb ligne|nombre de ligne", ql):
        return {"answer": f"« {target.get('dataset_name')} » contient {len(df):,} lignes.", "type": "text"}
    if re.search(r"combien de colonne|nb colonne|nombre de colonne", ql) or (
            re.search(r"\bcolonne", ql) and re.search(r"combien|nombre|liste", ql)):
        return {"answer": f"« {target.get('dataset_name')} » a {len(cols)} colonnes : {', '.join(cols)}", "type": "text"}

    # --- liste des colonnes / variables / schéma
    if re.search(r"\b(colonne|variable|champ|schéma|schema|structure)\b", ql) and not mentioned:
        return {"answer": f"Colonnes de « {target.get('dataset_name')} » : {', '.join(cols)}", "type": "text"}

    # --- agrégation
    agg = None
    if re.search(r"\b(somme|total|additionn)\b", ql): agg = "sum"
    elif re.search(r"\bmoyenne\b", ql): agg = "mean"
    elif re.search(r"\bmaximum|max\b", ql): agg = "max"
    elif re.search(r"\bminimum|min\b", ql): agg = "min"
    elif re.search(r"\bnombre de|combien de", ql): agg = "count"
    agg_label = {"sum": "Somme", "mean": "Moyenne", "max": "Maximum", "min": "Minimum", "count": "Nombre"}.get(agg)

    # --- colonne de regroupement (« par X »)
    group_col = None
    m = re.search(r"\bpar\s+([a-z0-9_ éàèêç\-]+)", ql)
    if m:
        cand = m.group(1).strip()
        for k, c in colmap.items():
            if k and (k in cand or cand in k):
                group_col = c
                break

    # --- colonne mesurée (mentionnée et numérique, sinon 1ère numérique)
    value_col = next((c for c in mentioned if c in numeric), None)
    if value_col is None and numeric:
        value_col = numeric[0]

    # --- top N
    topn = None
    m = re.search(r"\btop\s*(\d+)", ql)
    if m:
        topn = int(m.group(1))

    # --- exécution : groupement + agrégation
    if group_col and value_col and agg:
        g = df.groupby(group_col)[value_col].agg(agg).sort_values(ascending=False)
        rows = [[str(k), round(float(v), 2)] for k, v in g.items()]
        return {"answer": f"{agg_label} de « {value_col} » par « {group_col} » — dataset « {target.get('dataset_name')} » :",
                "type": "table", "table": {"columns": [group_col, f"{agg_label} {value_col}"], "rows": rows[:30]}}

    # --- agrégation simple
    if value_col and agg:
        try:
            v = getattr(df[value_col], agg)()
        except Exception:
            v = None
        if v is not None:
            return {"answer": f"{agg_label} de « {value_col} » = {round(float(v), 2)}", "type": "text"}

    # --- top N
    if topn and value_col:
        head = df.nlargest(topn, value_col)[[value_col] + ([group_col] if group_col else [])]
        rows = [[str(r[value_col]), str(r[group_col])] if group_col else [str(r[value_col])]
                for _, r in head.iterrows()]
        cols_out = [value_col] + ([group_col] if group_col else [])
        return {"answer": f"Top {topn} « {value_col} » :", "type": "table",
                "table": {"columns": cols_out, "rows": rows}}

    # --- aperçu du dataset par défaut
    head = df.head(5)
    rows = [[str(r.get(c)) for c in cols] for _, r in head.iterrows()]
    return {"answer": f"Dataset « {target.get('dataset_name')} » ({len(df):,} lignes × {len(cols)} colonnes). "
                     f"Aperçu (5 premières lignes) — posez par ex. « somme de {value_col or '…'} par {group_col or '…'} » :",
            "type": "table", "table": {"columns": cols, "rows": rows}}

# ---------- Chat LLM (DeepSeek, texte -> pandas + visualisation) ----------
def _llm_config():
    if T.CLOUD:
        key = os.environ.get("DEEPSEEK_API_KEY")
        return {"api_key": key, "model": os.environ.get("DEEPSEEK_MODEL", "deepseek-v4-pro"),
                "base_url": os.environ.get("DEEPSEEK_BASE_URL", "https://api.deepseek.com/chat/completions")} if key else None
    p = os.path.join(BASE, "llm_config.json")
    if os.path.exists(p):
        try:
            with open(p, encoding="utf-8") as f:
                return json.load(f)
        except Exception:
            return None
    return None

def _llm_call(messages, max_tokens=8000, temperature=0, timeout=120, model=None, effort=None, attempts=2):
    cfg = _llm_config()
    if not cfg or not cfg.get("api_key"):
        return None
    payload = {"model": model or cfg.get("model", "deepseek-v4-pro"), "messages": messages,
               "temperature": temperature, "max_tokens": max_tokens}
    if effort:
        payload["reasoning_effort"] = effort
    body = json.dumps(payload).encode("utf-8")
    for attempt in range(attempts):
        req = urllib.request.Request(cfg.get("base_url", "https://api.deepseek.com/chat/completions"),
                                     data=body, method="POST",
                                     headers={"Content-Type": "application/json",
                                              "Authorization": "Bearer " + cfg["api_key"]})
        try:
            r = urllib.request.urlopen(req, timeout=timeout)
            return json.loads(r.read().decode())["choices"][0]["message"]["content"]
        except Exception:
            if attempt == 0:
                time.sleep(2)
    return None

def _registry_datasets():
    """Jeux de données du registre ; liste vide (et trace dans les journaux) s'il est injoignable,
    pour que le chat réponde quand même sur le fichier en cours."""
    try:
        reg = T.registry()
        reg.ensure()
        return reg.latest_versions()
    except Exception as e:
        sys.stderr.write(f"[registre] indisponible pour le chat : {type(e).__name__} : {e}\n")
        return []


def _llm_context(datasets=None):
    datasets = _registry_datasets() if datasets is None else datasets
    out = []
    for d in datasets[:12]:
        df = _chat_load_work(d.get("work_path"))
        cols = []
        if df is not None:
            for c in df.columns:
                try:
                    vals = [str(v) for v in df[c].dropna().head(3).tolist()]
                except Exception:
                    vals = []
                cols.append({"name": str(c), "dtype": str(df[c].dtype), "exemples": vals})
        out.append({"name": d.get("dataset_name"), "id": d.get("dataset_id"),
                    "lignes": d.get("row_count"), "colonnes": cols})
    return out

def _work_file(path):
    """Fichier de travail (parquet/csv) d'une version du registre, ou None."""
    if not path:
        return None
    p = str(path)
    if os.path.isfile(p):
        return p
    if os.path.isdir(p):
        c = [os.path.join(p, f) for f in sorted(os.listdir(p)) if f.endswith((".parquet", ".csv"))]
        return c[0] if c else None
    for ext in (".parquet", ".csv"):
        if os.path.exists(p + ext):
            return p + ext
    return None


def _parse_llm_json(content):
    """Objet JSON renvoyé par l'IA (éventuellement entouré de ```), ou None."""
    text = (content or "").strip()
    if text.startswith("```"):
        text = re.sub(r"^```[a-zA-Z]*\s*", "", text)
        text = re.sub(r"\s*```$", "", text)
    try:
        return json.loads(text)
    except Exception:
        pass
    # Plusieurs objets à la suite (ex. un calcul ET une conclusion écrite sans en connaître le résultat) :
    # seul le PREMIER compte ; la suite, rédigée à l'aveugle, est ignorée.
    start = text.find("{")
    if start >= 0:
        try:
            obj, _end = json.JSONDecoder().raw_decode(text[start:])
            return obj
        except Exception:
            return None
    return None


@_lourd
def _run_ai_code(code, paths):
    try:
        return sandbox.run(code, paths)
    except sandbox.UnsafeCode as e:
        return {"error": f"calcul refusé par le contrôle de sécurité ({e})"}


INTERPRETATION_RULES = (
    "INTERPRÉTATION. Si la question demande un résumé, une interprétation, une explication ou « ce qu'on retient » "
    "de l'analyse, réponds DIRECTEMENT dans \"answer\" à partir du RÉSULTAT DE L'ANALYSE ci-dessus, sans \"code\" :\n"
    "- langage simple pour un responsable non technique ; 4 à 8 puces courtes (« - ») ou 2-3 paragraphes ;\n"
    "- chaque constat appuyé sur un chiffre du résultat (arrondi simplement) ;\n"
    "- lecture CRITIQUE : une « valeur aberrante » sur un identifiant ou un code (code_site, id, numéro…) n'a pas de "
    "sens statistique, dis-le ; des bornes négatives pour une grandeur positive signalent une distribution très "
    "asymétrique ; en santé, une valeur extrême peut être un vrai cas clinique et non une erreur ;\n"
    "- sépare les faits (chiffres) des hypothèses (« probablement », « à vérifier ») ;\n"
    "- termine par 1 à 3 suites concrètes (vérifier, croiser, approfondir).\n"
)


def _analysis_for_chat(cur_path):
    """Résultat de l'analyse du fichier en cours, recalculé s'il n'est plus en mémoire (redémarrage)."""
    if not cur_path:
        return None
    st = T.state()
    if st.get("last_result") and st.get("result_path") == cur_path:
        return st["last_result"]
    res = analyze(cur_path)
    if "error" in res:
        return None
    st["last_result"], st["result_path"] = res, cur_path
    return res


def _analysis_context(res, ui_context):
    summary = studio._summary(res)
    summary["associations_categories"] = ((res.get("bivariate") or {}).get("cramer") or [])[:6]
    if res.get("controles") is not None:  # rapport remis à plat : écarts total / sexe / âge
        summary["controles_coherence"] = {"nombre_d_ecarts": len(res["controles"]), "exemples": res["controles"][:15]}
    out = {"resultat_analyse": summary}
    if isinstance(ui_context, dict):  # ce que seul l'écran connaît : corrections appliquées, structure, points clés
        text = json.dumps(ui_context, ensure_ascii=False, default=str)
        out["ecran"] = ui_context if len(text) <= 4000 else text[:4000]
    return json.dumps(out, ensure_ascii=False, default=str)


CHAT_MAX_CALCULS = 4

CHAT_METHODE = (
    "MÉTHODE DE TRAVAIL — tu es un agent d'analyse. Tu réponds à TOUTE question sur les données : description, "
    "filtres, comparaisons entre groupes, tendances, répartitions, croisements, tests statistiques, modèles simples, "
    "qualité des données, interprétation de l'analyse.\n"
    "À chaque tour, réponds UNIQUEMENT par UN SEUL objet JSON valide, sans texte autour. Ne mets JAMAIS un calcul "
    "et une réponse dans le même tour : tant que tu n'as pas reçu le résultat, tu ne le connais pas.\n"
    '- pour calculer : {"action":"calcul","pourquoi":"ce que ce calcul vérifie, en une phrase","code":"..."}\n'
    '- pour conclure : {"action":"reponse","answer":"ta réponse finale","montrer_tableau":true}\n'
    f"Après chaque calcul tu reçois son résultat RÉEL. Enchaîne d'autres calculs si nécessaire (au plus {CHAT_MAX_CALCULS}), "
    "puis conclus. Si la question porte seulement sur l'analyse déjà affichée (résumé, interprétation), conclus sans calcul.\n"
    "CODE : Python avec pandas/numpy. Les données sont dans le dict `data` (clé = nom du dataset, valeur = DataFrame). "
    "pd, np, math et stats (= scipy.stats : ttest_ind, mannwhitneyu, chi2_contingency, pearsonr, spearmanr, linregress, "
    "f_oneway, kruskal, shapiro…) sont DÉJÀ disponibles : n'écris aucune ligne import. "
    "Mets ce que tu veux lire dans la variable `result` (DataFrame, Series ou valeur) et/ou affiche-le avec print(). "
    "Un tableau renvoyé est limité à 50 lignes : agrège ou trie avant. "
    "Écris les noms de colonnes ENTRE GUILLEMETS DOUBLES, exactement comme listés : df[\"Nom d'utilisateur\"].\n"
    "GRAPHIQUE : dans le code, assigne une variable `chart` = {\"type\":\"...\",\"title\":\"...\",\"labels\":[...],"
    "\"values\":[[...]],\"seriesNames\":[...]} (liste de dicts pour plusieurs graphiques). Types : bar (catégories), "
    "line ou area (évolution dans le temps, trie par période), pie/doughnut (parts), scatter (labels = x, values = [y]), "
    "histogram (labels = intervalles, values = effectifs), box (values[i] = [min, Q1, médiane, Q3, max] par catégorie). "
    "Aucune bibliothèque de graphique (pas de matplotlib). Le graphique du dernier calcul qui en définit un est affiché.\n"
    "RÉPONSE FINALE : commence par la réponse directe à la question. Chiffres UNIQUEMENT issus des résultats reçus ou "
    "de l'analyse fournie, jamais inventés. Cite la méthode quand c'est un test (ex. « test de Student, p = 0,03 ») et "
    "les limites (petits effectifs, valeurs manquantes, corrélation ≠ causalité). Court pour une question simple, "
    "structuré en puces pour une analyse. montrer_tableau = true si le dernier tableau calculé aide à répondre. "
    "Si un calcul échoue, corrige-le au tour suivant. Si les données ne permettent pas de répondre, dis-le.\n"
)


def _chat_llm(messages, max_tokens=8000):
    """Modèle rapide pour l'agent (plusieurs allers-retours par question), réflexion réduite en secours."""
    cfg = _llm_config()
    model = os.environ.get("KLDV_CHAT_MODEL") or (cfg or {}).get("chat_model") or "deepseek-flash"
    return (_llm_call(messages, max_tokens, model=model, effort="high", timeout=90, attempts=1)
            or _llm_call(messages, max_tokens, model=model, effort="low", timeout=60, attempts=1))


def _format_calc(out):
    """Résultat d'un calcul, tel que l'IA le relit (texte compact)."""
    if out.get("error"):
        text = f"ERREUR : {out['error']}"
    else:
        res = out.get("result") or {}
        if res.get("kind") == "table":
            rows = [" | ".join(res["columns"])] + [" | ".join(r) for r in res["rows"][:50]]
            text = f"Tableau ({len(res['rows'])} lignes affichées) :\n" + "\n".join(rows)
        elif res.get("kind") == "scalar":
            text = "Valeur : " + str(res.get("value"))
        else:
            text = "(aucune variable result)"
        ch = out.get("chart")
        if ch:
            titles = [c.get("title", "") for c in (ch if isinstance(ch, list) else [ch]) if isinstance(c, dict)]
            text += "\nGraphique préparé : " + ", ".join(titles)
    if out.get("sortie"):
        text += "\nSortie de print() :\n" + out["sortie"]
    return text[:6000]


def _charts_list(chart):
    if isinstance(chart, list):
        return [c for c in chart if isinstance(c, dict) and c.get("labels") and c.get("values")]
    if isinstance(chart, dict) and chart.get("labels") and chart.get("values"):
        return [chart]
    return []


def llm_chat(question, ui_context=None, history=None, progress=None, linked=None):
    """Agent d'analyse : l'IA enchaîne des calculs (exécutés dans le bac à sable), lit leurs résultats
    réels, puis rédige sa réponse. `progress(texte)` reçoit l'étape en cours (affichage en direct).
    `linked` : [{path, name}] fichiers liés au fichier en cours (autres feuilles du classeur, fichiers croisés)."""
    progress = progress or (lambda _t: None)
    datasets = _registry_datasets()
    context = _llm_context(datasets)
    current_df, current_name = None, None
    cur_path = T.state()["current_path"]
    try:
        analysis = _analysis_for_chat(cur_path)
    except Exception as e:  # l'interprétation est un plus : le chat répond quand même sans elle
        sys.stderr.write(f"[chat] analyse indisponible : {type(e).__name__} : {e}\n")
        analysis = None
    if cur_path:
        current_df, _err = _read_df(cur_path)
        if current_df is not None:
            current_name = source_label(cur_path)
            context.insert(0, {"name": current_name, "fichier_en_cours_d_analyse": True, "lignes": int(len(current_df)),
                               "colonnes": [{"name": str(c), "dtype": str(current_df[c].dtype),
                                             "exemples": [str(v) for v in current_df[c].dropna().head(3).tolist()]}
                                            for c in current_df.columns]})
        if current_df is not None and _is_report(current_df):
            context[0]["regle_rapport"] = (
                "Rapport remis à plat : une ligne = indicateur × niveau × sexe × tranche_age × site. La colonne niveau "
                "vaut total, sexe ou age. Ne JAMAIS additionner deux niveaux : total d'un indicateur → niveau == 'total' ; "
                "par sexe → niveau == 'sexe' ; par âge (et sexe) → niveau == 'age'. Ne pas additionner des indicateurs différents.")
    linked_paths = {}
    for item in (linked or [])[:MAX_ARCHITECT_FILES]:  # autres feuilles / fichiers croisés, aussi calculables
        lp = str((item or {}).get("path") or "")
        if not lp or lp == cur_path or lp.startswith(("http://", "https://")):
            continue
        ldf, _err = _read_df(lp)  # vérifie aussi que le fichier appartient à l'organisation
        if ldf is None:
            continue
        lname = str(item.get("name") or source_label(lp))
        # Copie parquet pour le bac à sable (qui ne lit que parquet/CSV), comme le fichier en cours
        tmp = os.path.join(T.current().dir("tmp"), f"lie_{T.current().key[:12]}_{len(linked_paths)}.parquet")
        _save_table(ldf, tmp)
        linked_paths[lname] = tmp
        context.insert(1, {"name": lname, "fichier_lie_au_fichier_en_cours": True, "lignes": int(len(ldf)),
                           "colonnes": [{"name": str(c), "dtype": str(ldf[c].dtype),
                                         "exemples": [str(v) for v in ldf[c].dropna().head(3).tolist()]}
                                        for c in ldf.columns[:80]]})
    system = (
        "Tu es Hermes, l'analyste de données de KLDataViz. Réponds en français, de façon fiable.\n"
        + studio.chat_preamble() +
        "Sauf mention contraire, la question porte sur le fichier en cours d'analyse (premier dataset de la liste).\n"
        "Datasets disponibles (nom, nombre de lignes, colonnes avec type et exemples) :\n"
        + json.dumps(context, ensure_ascii=False, default=str) + "\n\n"
        + ("RÉSULTAT DE L'ANALYSE affichée à l'utilisateur (statistiques déjà calculées sur le fichier en cours) :\n"
           + _analysis_context(analysis, ui_context) + "\n\n" + INTERPRETATION_RULES + "\n" if analysis else "")
        + CHAT_METHODE
    )
    messages = [{"role": "system", "content": system}]
    for h in (history or [])[-6:]:  # mémoire courte de la conversation
        if isinstance(h, dict) and h.get("role") in ("user", "assistant") and h.get("text"):
            messages.append({"role": h["role"], "content": str(h["text"])[:1500]})
    messages.append({"role": "user", "content": question})

    paths = None

    def data_paths():  # préparés une seule fois, au premier calcul
        nonlocal paths
        if paths is None:
            paths = {}
            for _d in datasets:
                _p = _work_file(_d.get("work_path"))
                if _p and T.current().owns(_p):
                    paths[_d.get("dataset_name") or _d.get("dataset_id")] = _p
            if current_df is not None:
                _cur = os.path.join(T.current().dir("tmp"), f"courant_{T.current().key[:12]}.parquet")
                try:
                    current_df.to_parquet(_cur)
                except Exception:
                    current_df.astype(str).to_parquet(_cur)
                paths[current_name] = _cur
            paths.update(linked_paths)
        return paths

    etapes, last_table, last_chart = [], None, None
    progress("Hermes réfléchit à la question…")
    for _tour in range(CHAT_MAX_CALCULS + 2):
        content = _chat_llm(messages)
        if not content:
            if not etapes:
                return None  # IA indisponible : repli par règles côté route
            break
        data = _parse_llm_json(content)
        if not isinstance(data, dict):
            return {"answer": content, "type": "text", "table": None, "charts": _charts_list(last_chart), "etapes": etapes}
        action = data.get("action") or ("calcul" if data.get("code") else "reponse")
        if action == "calcul" and data.get("code") and len(etapes) < CHAT_MAX_CALCULS:
            n = len(etapes) + 1
            pourquoi = str(data.get("pourquoi") or "calcul")[:200]
            progress(f"Calcul {n} : {pourquoi}")
            out = _run_ai_code(str(data["code"]), data_paths())
            feedback = _format_calc(out)
            etapes.append({"pourquoi": pourquoi, "code": str(data["code"])[:4000], "ok": not out.get("error"),
                           "resultat": feedback[:1500]})
            if not out.get("error"):
                last_chart = out.get("chart") or last_chart
                res = out.get("result") or {}
                if res.get("kind") == "table":
                    last_table = {"columns": res["columns"], "rows": res["rows"]}
            suite = ("Limite de calculs atteinte : conclus maintenant avec action \"reponse\"."
                     if len(etapes) >= CHAT_MAX_CALCULS else
                     "Fais un autre calcul si c'est nécessaire, sinon conclus avec action \"reponse\".")
            messages += [{"role": "assistant", "content": content},
                         {"role": "user", "content": f"Résultat réel du calcul {n} :\n{feedback}\n\n{suite}"}]
            progress("Hermes lit le résultat…")
            continue
        answer = str(data.get("answer") or data.get("reponse") or "").strip()
        if action == "calcul" and not answer:  # limite atteinte mais l'IA voulait encore calculer
            messages += [{"role": "assistant", "content": content},
                         {"role": "user", "content": "Plus de calcul possible : conclus avec action \"reponse\"."}]
            continue
        show = data.get("montrer_tableau", True) is not False
        return {"answer": answer or "(réponse vide)", "type": "table" if (last_table and show) else "text",
                "table": last_table if show else None, "charts": _charts_list(last_chart), "etapes": etapes}
    last = etapes[-1]["resultat"] if etapes else ""
    return {"answer": "Je n'ai pas pu conclure dans le temps imparti. Dernier résultat obtenu :\n" + last,
            "type": "table" if last_table else "text", "table": last_table, "charts": _charts_list(last_chart),
            "etapes": etapes}


def _json_safe(obj):
    import math
    if obj is None or isinstance(obj, (str, bool, int)):
        return obj
    if isinstance(obj, float):
        return None if (math.isnan(obj) or math.isinf(obj)) else obj
    if hasattr(obj, "dtype") and hasattr(obj, "item"):  # scalaire numpy
        try:
            return _json_safe(obj.item())
        except Exception:
            return str(obj)
    if isinstance(obj, dict):
        return {k: _json_safe(v) for k, v in obj.items()}
    if isinstance(obj, (list, tuple)):
        return [_json_safe(v) for v in obj]
    return obj

def _architect_df(file_path):
    """Lit la source pour l'architecte : détection d'encodage et de séparateur (« ; » des exports FR/DHIS2)."""
    if file_path.startswith(("http://", "https://", "dhis2://")):
        df, err = _read_df(file_path)
        return df, {"fichier": source_label(file_path)}, err
    if T.CLOUD and not T.current().owns(file_path):
        return None, None, "Fichier introuvable."
    if not os.path.isfile(file_path):
        return None, None, f"Fichier introuvable : {file_path}"
    try:
        df, info = data_architect.load_table(file_path)
    except Exception as e:
        return None, None, f"Lecture impossible : {e}"
    if os.path.exists(file_path + AGREGE_SUFFIX):  # fichier dérivé d'une source agrégée : comme _read_df
        df.attrs["agrege"] = True
    return df, info, None


MAX_ARCHITECT_FILES = 6


def _own_excel(file_path):
    """Erreur si `file_path` n'est pas un classeur Excel local de l'organisation, sinon None."""
    if T.CLOUD and not T.current().owns(file_path):
        return "Fichier introuvable."
    if not os.path.isfile(file_path):
        return f"Fichier introuvable : {file_path}"
    if os.path.splitext(file_path)[1].lower() not in data_architect.EXCEL_EXT:
        return "Ce fichier n'est pas un classeur Excel."
    return None


def sheets_info(file_path):
    """Feuilles d'un classeur, avec leurs dimensions et la feuille recommandée pour l'analyse."""
    err = _own_excel(file_path)
    if err:
        return {"error": err}
    try:
        return {"feuilles": data_architect.excel_sheets(file_path)}
    except Exception as e:
        return {"error": f"Lecture du classeur impossible : {e}"}


def _save_table(df, dest):
    """Parquet (types conservés) ; colonnes de types mélangés converties en texte si nécessaire."""
    df = df.copy()
    df.attrs = {}  # pandas les enregistre dans le parquet : la feuille d'origine ne doit pas suivre le fichier
    df.columns = [str(c) for c in df.columns]
    try:
        df.to_parquet(dest, index=False)
    except Exception:
        for c in df.columns[df.dtypes == object]:
            df[c] = df[c].map(lambda v: v if v is None or (isinstance(v, float) and v != v) else str(v))
        df.to_parquet(dest, index=False)


MAX_STACK_SHEETS = 300


def sheets_extract(file_path, sheets, stack=False):
    """Une feuille → un fichier de travail : toutes les étapes (analyse, nettoyage, architecte, chat)
    s'appliquent ensuite à elle seule, sans ambiguïté sur la feuille lue.
    `stack` : les feuilles (même structure) sont empilées en un seul fichier, colonne « feuille » en tête."""
    err = _own_excel(file_path)
    if err:
        return {"error": err}
    names = pd.ExcelFile(file_path).sheet_names
    wanted = [str(x) for x in (sheets or []) if str(x) in names]
    wanted = wanted[:MAX_STACK_SHEETS if stack else MAX_ARCHITECT_FILES]
    if not wanted:
        return {"error": "Aucune feuille valide demandée."}
    stem = os.path.splitext(os.path.basename(file_path))[0]
    if stack:
        try:
            meta = {f["nom"]: f for f in data_architect.excel_sheets(file_path)}
            if all(meta.get(n, {}).get("rapport") for n in wanted):  # rapports : chacun remis à plat, puis empilés
                frames = []
                for n in wanted:
                    one, err = _read_df(file_path, sheet=n)
                    if err:
                        return {"error": f"Feuille « {n} » : {err}"}
                    frames.append(one.assign(feuille=n)[["feuille"] + list(one.columns)])
                df = pd.concat(frames, ignore_index=True)
            else:
                df = _tidy(data_architect.stack_sheets(file_path, wanted))
        except Exception as e:
            return {"error": f"Empilement impossible : {e}"}
        dest = os.path.join(T.uploads_dir(), _safe_name(f"{stem} - {len(wanted)} feuilles empilees.parquet"))
        _save_table(df, dest)
        return {"fichiers": [{"path": dest, "name": f"{stem} — {len(wanted)} feuilles empilées", "feuille": None,
                              "feuilles": wanted, "lignes": int(len(df)), "colonnes": int(df.shape[1])}]}
    out = []
    for name in wanted:
        df, err = _read_df(file_path, sheet=name)
        if err:
            return {"error": f"Feuille « {name} » : {err}"}
        dest = os.path.join(T.uploads_dir(), _safe_name(f"{stem} - {name}.parquet"))
        _save_table(df, dest)
        out.append({"path": dest, "name": f"{stem} — {name}", "feuille": name,
                    "lignes": int(len(df)), "colonnes": int(df.shape[1])})
    return {"fichiers": out}


SYSTEM_CLASSEUR = """Tu es Hermes, analyste de données. On te donne plusieurs feuilles d'un même classeur Excel avec
des FAITS MESURÉS : le profil de chaque feuille, les formules qui relient les feuilles (qui calcule à partir de qui,
avec quelles fonctions), les valeurs communes entre colonnes (jointures possibles) et les feuilles de même structure.
Explique la logique du classeur puis propose des analyses. Appuie-toi uniquement sur ces faits ; un lien supposé
est marqué « hypothèse ». Écris en français clair, pour un responsable métier.
Réponds par UN SEUL objet JSON :
{"resume": "2 à 4 phrases : à quoi sert ce classeur et comment les feuilles s'enchaînent",
 "flux": "chaîne lisible, ex. « CD-IST, PEC, PTME → Nbre de problèmes par site → Feuil2 → DASHBOARD »",
 "roles": [{"feuille": "nom exact", "role": "données brutes|référentiel|paramètres|calcul intermédiaire|synthèse|tableau de bord|autre",
            "explication": "ce que contient la feuille et à quoi elle sert"}],
 "liens": [{"de": "feuille", "vers": "feuille", "nature": "formule|clé commune|même structure|hypothèse",
            "explication": "ce que fait le lien, en clair", "preuve": "le fait mesuré qui le montre"}],
 "points_attention": ["incohérence, risque de double comptage, plage de formule trop courte, feuille isolée…"],
 "analyses": [{"titre": "titre court", "objectif": "ce qu'on apprend", "feuilles": ["noms exacts"],
               "question": "question précise et autonome pour l'agent d'analyse, qui calculera lui-même sur les feuilles : cite les feuilles et les colonnes exactes"}]}
Propose 3 à 6 analyses concrètes et utiles au métier (comparaisons, évolutions, rapprochements et contrôles de
cohérence entre feuilles), de la plus utile à la moins utile."""


def _classeur_fallback(names, links, rels, groups):
    """Lecture par règles quand l'IA ne répond pas : rôles déduits du sens des formules."""
    uses = {n: {l["utilise"] for l in links if l["feuille"] == n} for n in names}
    used = {n: {l["feuille"] for l in links if l["utilise"] == n} for n in names}
    roles = []
    for n in names:
        role = ("données brutes" if used[n] and not uses[n] else "synthèse" if uses[n] and not used[n]
                else "calcul intermédiaire" if uses[n] else "autre")
        roles.append({"feuille": n, "role": role, "explication":
                      (f"Utilisée par {', '.join(sorted(used[n]))}. " if used[n] else "")
                      + (f"Calcule à partir de {', '.join(sorted(uses[n]))}." if uses[n] else "")})
    liens = [{"de": l["utilise"], "vers": l["feuille"], "nature": "formule",
              "explication": f"{l['feuille']} reprend ou calcule des valeurs de {l['utilise']}",
              "preuve": f"{l['formules']} formule(s), ex. {l['exemple']}"} for l in links]
    liens += [{"de": r["gauche"].split(" :: ")[0], "vers": r["droite"].split(" :: ")[0], "nature": "clé commune",
               "explication": f"{r['gauche']} et {r['droite']} partagent des valeurs ({r['cardinalite']})",
               "preuve": f"{r['valeurs_communes']} valeurs communes"} for r in rels[:6]]
    analyses = [{"titre": f"Rapprocher {r['gauche'].split(' :: ')[0]} et {r['droite'].split(' :: ')[0]}",
                 "objectif": "Vérifier que les deux feuilles concordent et repérer les écarts",
                 "feuilles": [r["gauche"].split(" :: ")[0], r["droite"].split(" :: ")[0]],
                 "question": f"Rapproche {r['gauche']} et {r['droite']} : combien de valeurs sont présentes "
                             f"dans une feuille et absentes de l'autre, et lesquelles ?"} for r in rels[:3]]
    for g in groups[:2]:
        analyses.append({"titre": f"Comparer les {len(g)} feuilles de même structure",
                         "objectif": "Voir l'évolution ou les écarts d'une feuille à l'autre", "feuilles": g,
                         "question": f"Compare les feuilles {', '.join(g)} (mêmes colonnes) : totaux par feuille "
                                     "et principales différences."})
    ordre = list(dict.fromkeys([l["utilise"] for l in reversed(links)] + [l["feuille"] for l in reversed(links)]))
    return {"resume": "Lecture par règles : l'interprétation par l'IA n'est pas disponible.",
            "flux": " → ".join(ordre), "roles": roles, "liens": liens, "points_attention": [], "analyses": analyses}


def workbook_understand(file_path, sheets, progress=lambda _t: None):
    """Logique d'un classeur : rôle de chaque feuille cochée, liens entre elles (formules, clés communes,
    même structure) et analyses proposées. Faits mesurés par le moteur, interprétation par l'IA."""
    err = _own_excel(file_path)
    if err:
        return {"error": err}
    wanted = {str(x) for x in sheets or []}
    names = [n for n in pd.ExcelFile(file_path).sheet_names if n in wanted][:MAX_ARCHITECT_FILES]
    if len(names) < 2:
        return {"error": "Cochez au moins deux feuilles."}
    progress("Lecture des feuilles")
    frames, profils = {}, {}
    for n in names:
        df, e = _read_df(file_path, sheet=n)
        if e:
            return {"error": f"Feuille « {n} » : {e}"}
        frames[n] = df
        try:
            prof = architect._compact_profile(data_architect.profile(df, {"fichier": n}), df) if architect else {}
        except Exception:
            prof = {}
        if len(json.dumps(prof, ensure_ascii=False, default=str)) > 8000:  # trop de colonnes : l'essentiel
            prof = {"lignes": int(len(df)),
                    "colonnes": {str(c): [str(df[c].dtype)] + [str(v) for v in df[c].dropna().head(2)]
                                 for c in df.columns[:60]}}
        profils[n] = prof
    progress("Recherche des formules entre feuilles")
    try:
        f = data_architect.workbook_links(file_path)
    except Exception as e:
        f = {"disponible": False, "liens": [], "erreur": str(e)}
    links = [l for l in f.get("liens", []) if l["feuille"] in names or l["utilise"] in names]
    progress("Recherche des colonnes communes")
    rels = data_architect.find_relations(frames)
    meta = {x["nom"]: x for x in data_architect.excel_sheets(file_path)}
    groups = {}
    for n in names:
        if meta.get(n, {}).get("groupe"):
            groups.setdefault(meta[n]["groupe"], []).append(n)
    groups = [g for g in groups.values() if len(g) > 1]
    faits = {"feuilles": {n: {"masquee": meta.get(n, {}).get("masquee"), "remarques": meta.get(n, {}).get("notes"),
                              "profil": profils[n]} for n in names},
             "formules_entre_feuilles": links[:30],
             "formules_internes": {n: f.get("calculs_internes", {}).get(n, 0) for n in names},
             "formules_lisibles": f.get("disponible", False),
             "colonnes_communes": rels[:15], "feuilles_meme_structure": groups}
    result, ia, cfg = None, False, _llm_config()
    if cfg:
        progress("Hermes interprète la logique du classeur")
        model = os.environ.get("KLDV_ARCHITECT_MODEL") or cfg.get("architect_model") or "deepseek-flash"
        msgs = [{"role": "system", "content": SYSTEM_CLASSEUR},
                {"role": "user", "content": "Classeur : " + os.path.basename(file_path) + "\nFAITS :\n"
                 + json.dumps(faits, ensure_ascii=False, default=str)[:60000]}]
        for effort, timeout in (("high", 120), ("low", 90)):
            result = _parse_llm_json(_llm_call(msgs, 6000, model=model, effort=effort, timeout=timeout, attempts=1))
            if isinstance(result, dict) and result.get("analyses"):
                ia = True
                break
    if not ia:
        result = _classeur_fallback(names, links, rels, groups)
    for a in result.get("analyses") or []:  # feuilles citées : seulement celles qui ont été étudiées
        a["feuilles"] = [x for x in a.get("feuilles") or [] if x in names] or names
    return {"comprehension": result, "ia": ia, "ia_disponible": bool(cfg), "feuilles": names,
            "faits": {"formules_entre_feuilles": links[:30], "colonnes_communes": rels[:15],
                      "feuilles_meme_structure": groups, "formules_lisibles": f.get("disponible", False)}}


def _architect_paths(p):
    """Chemins envoyés par l'interface : `file_paths` (plusieurs fichiers) ou `file_path` (un seul)."""
    paths = p.get("file_paths")
    return [str(x) for x in paths] if isinstance(paths, list) else [str(p.get("file_path") or "")]


def _architect_tables(paths):
    """{nom de fichier unique: (DataFrame, info)} pour une liste de chemins, ou (None, erreur)."""
    paths = [str(p) for p in (paths or []) if p][:MAX_ARCHITECT_FILES]
    if not paths:
        return None, "Aucun fichier."
    tables = {}
    for p in paths:
        df, info, err = _architect_df(p)
        if err:
            return None, f"{os.path.basename(p)} : {err}"
        name = (info or {}).get("fichier") or os.path.basename(p)
        while name in tables:
            name = "_" + name
        tables[name] = (df, info)
    return tables, None


def architect_propose(paths):
    if architect is None:
        return {"error": f"Agent architecte indisponible : {ARCHITECT_ERROR}"}
    tables, err = _architect_tables(paths)
    if err:
        return {"error": err}
    cfg = _llm_config()
    # Modèle rapide par défaut : deepseek-flash répond en ~1 min contre 3 à 7 min pour deepseek-v4-pro,
    # avec une qualité équivalente sur cette tâche (le moteur fournit déjà les faits mesurés).
    # Si l'API traîne, on réessaie une fois avec moins de réflexion (~30 s), puis repli par règles.
    model = os.environ.get("KLDV_ARCHITECT_MODEL") or (cfg or {}).get("architect_model") or "deepseek-flash"

    def llm(messages, max_tokens):
        return (_llm_call(messages, max_tokens, model=model, effort="high", timeout=120, attempts=1)
                or _llm_call(messages, max_tokens, model=model, effort="low", timeout=90, attempts=1))
    if len(tables) == 1:
        df, info = next(iter(tables.values()))
        return _json_safe(architect.propose(df, info, llm if cfg else None))
    return _json_safe(architect.propose_many(tables, llm if cfg else None))


# L'analyse IA dure plusieurs minutes : elle tourne en tâche de fond et le navigateur
# interroge son état, pour ne pas dépasser le délai des proxys (Render coupe vers 100 s).
_arch_jobs, _arch_lock = {}, threading.Lock()


def start_job(fn):
    """Lance fn(progress) en tâche de fond pour l'utilisateur courant ; progress(texte) publie l'étape en cours."""
    ctx = T.current()
    job_id = secrets.token_hex(8)

    def progress(text):
        with _arch_lock:
            if job_id in _arch_jobs:
                _arch_jobs[job_id]["etape"] = str(text)[:200]

    def run():
        with T.use(ctx):
            try:
                result = _json_safe(fn(progress))
            except Exception as e:
                sys.stderr.write(f"[tâche] {type(e).__name__} : {e}\n" + "".join(traceback.format_exception(e)[-4:]))
                result = {"error": f"Traitement impossible : {e}"}
        with _arch_lock:
            _arch_jobs[job_id].update(status="done", result=result)

    with _arch_lock:
        for k in [k for k, j in _arch_jobs.items() if time.time() - j["t"] > 3600]:
            del _arch_jobs[k]
        _arch_jobs[job_id] = {"key": ctx.key, "status": "running", "result": None, "etape": "", "t": time.time()}
    threading.Thread(target=run, daemon=True).start()
    return {"job": job_id}


def job_status(job_id):
    with _arch_lock:
        j = _arch_jobs.get(job_id)
        if not j or j["key"] != T.current().key:  # une tâche n'est visible que par son auteur
            return {"error": "Tâche introuvable ou expirée."}
        return {"status": j["status"], "result": j["result"], "etape": j.get("etape", "")}


def architect_start(paths):
    return start_job(lambda progress: architect_propose(paths))


architect_job = job_status


def _period_series(s):
    """Colonne de période sous forme de dates (date réelle, « 2026-01 », « Avril 2025 », « T1 2025 »), sinon None."""
    if pd.api.types.is_datetime64_any_dtype(s):
        return s
    if pd.api.types.is_numeric_dtype(s):
        return None
    sample = s.dropna().astype(str).head(40)
    if len(sample) < 2:
        return None
    try:
        parsed = to_dates(sample)
        if parsed.notna().mean() >= 0.9:
            return to_dates(s)
    except Exception:
        pass
    if sample.map(_fr_period).notna().mean() >= 0.9:
        return pd.to_datetime(s.map(_fr_period))
    return None


def _rule_items(rapport):
    """Corrections proposées par les règles, avec les mêmes identifiants que l'interface (étape Vérifier)."""
    items = []
    if rapport.get("doublons_supprimes"):
        items.append({"id": "doublons", "correction": f"supprimer {rapport['doublons_supprimes']} ligne(s) en double"})
    if rapport.get("colonnes_vides"):
        items.append({"id": "vides", "correction": "supprimer les colonnes vides : " + ", ".join(rapport["colonnes_vides"][:8])})
    for k in rapport.get("encodage_corrige") or {}:
        items.append({"id": f"encodage:{k}", "correction": f"réparer l'encodage du texte de « {k} »"})
    for k, v in (rapport.get("types_corriges") or {}).items():
        items.append({"id": f"type:{k}", "correction": f"convertir « {k} » de {v.get('avant')} en {v.get('apres')}"})
    for im in rapport.get("imputations") or []:
        items.append({"id": f"imput:{im['colonne']}",
                      "correction": f"compléter {im['n']} vide(s) de « {im['colonne']} » par la {im.get('methode')} ({im.get('valeur')})"})
    return items[:40]


def hermes_clean(file_path, progress=lambda _t: None):
    """Avis de Hermes sur les corrections des règles et corrections supplémentaires vérifiées."""
    import hermes_clean as HCL
    df, err = _read_df(file_path)
    if err:
        return {"error": err}
    _, rapport = clean_dataframe(df)
    cfg = _llm_config()
    model = os.environ.get("KLDV_CLEAN_MODEL") or (cfg or {}).get("clean_model") or "deepseek-flash"

    def llm(messages, max_tokens):
        return (_llm_call(messages, max_tokens, model=model, effort="high", timeout=100, attempts=1)
                or _llm_call(messages, max_tokens, model=model, effort="low", timeout=70, attempts=1))
    return HCL.run(df, _rule_items(rapport), source_label(file_path), llm if cfg else None, _parse_llm_json, progress)


def _registry_meta_dir():
    return T.current().dir(os.path.join("storage", "registry_meta"))


def schedule_commentary(job, result):
    """Rapport planifié : ce qui a changé depuis l'envoi précédent, rédigé par Hermes (règles sans IA ou hors quota)."""
    import hermes_report as HR
    cfg = _llm_config()
    llm = None
    if cfg:
        try:
            if T.CLOUD:
                cloud.consume(T.current(), "ai")
            model = os.environ.get("KLDV_REPORT_MODEL") or cfg.get("report_model") or "deepseek-flash"
            llm = lambda m, n: (_llm_call(m, n, model=model, effort="high", timeout=100, attempts=1)
                                or _llm_call(m, n, model=model, effort="low", timeout=70, attempts=1))
        except Exception:  # quota atteint : commentaire par règles
            llm = None
    return HR.run(job["id"], result, T.current().dir(os.path.join("storage", "snapshots")), llm, _parse_llm_json)


def hermes_charts(file_path, progress=lambda _t: None, comprehension=None):
    """Graphiques choisis et commentés par Hermes pour le fichier analysé (repli : graphiques automatiques)."""
    import hermes_charts as HC
    df, err = _read_df(file_path)
    if err:
        return {"error": err}
    cumul = _cumul_rows(df)
    if cumul["lignes"]:  # les lignes « Total » ne sont pas des catégories
        txt = [c for c in df.columns if not pd.api.types.is_numeric_dtype(df[c])]
        df = df[~df[txt].astype(str).apply(lambda s: s.str.match(TOTAL_LABEL)).any(axis=1)]
    base = _report_charts(df) if _is_report(df) else _smart_charts(df)
    cfg = _llm_config()
    model = os.environ.get("KLDV_CHARTS_MODEL") or (cfg or {}).get("charts_model") or "deepseek-flash"

    def llm(messages, max_tokens):
        return (_llm_call(messages, max_tokens, model=model, effort="high", timeout=100, attempts=1)
                or _llm_call(messages, max_tokens, model=model, effort="low", timeout=70, attempts=1))
    return HC.run(df, source_label(file_path), base, llm if cfg else None, _parse_llm_json, _period_series, progress,
                  comprehension if isinstance(comprehension, dict) else None)


def architect_apply(paths, spec):
    if architect is None:
        return {"error": f"Agent architecte indisponible : {ARCHITECT_ERROR}"}
    if not isinstance(spec, dict) or not spec:
        return {"error": "Aucune structure à générer."}
    tables, err = _architect_tables(paths)
    if err:
        return {"error": err}
    try:
        if len(tables) == 1:
            (name, (df, _)), = tables.items()
            r = architect.apply(df, spec, T.exports_dir(), name, views_dir=T.cleaned_dir())
        else:
            r = architect.apply_many(tables, spec, T.exports_dir(), views_dir=T.cleaned_dir())
    except Exception as e:
        return {"error": f"Restructuration impossible : {e}"}
    return {**r, "url": download_url(r["file"])}


# ---------------------------------------------------------------- serveur HTTP
SECRET = (os.environ.get("KLDV_SECRET") or secrets.token_hex(32)).encode()  # signe les liens de téléchargement
# En ligne, l'envoi arrive en base64 dans du JSON puis est décodé : ~3 fois sa taille en mémoire
MAX_UPLOAD = int(os.environ.get("KLDV_MAX_UPLOAD_MB", "25" if T.CLOUD else "50")) * 1024 * 1024
MAX_JSON = 2 * 1024 * 1024
ALLOWED_EXT = (".csv", ".xlsx", ".xls", ".xlsm", ".json", ".parquet")
EXPORTERS = {"pdf": export_pdf, "excel": export_excel, "html": export_html}
STATIC = {"/": ("index.html", "text/html"), "/index.html": ("index.html", "text/html"),
          "/login": ("login.html", "text/html"), "/logo.png": ("logo.png", "image/png"),
          "/favicon.png": ("logo.png", "image/png"), "/chart.umd.min.js": ("chart.umd.min.js", "application/javascript")}
if T.CLOUD:  # en ligne : page d'accueil publique, application sur /app, administration sur /admin
    STATIC.update({"/": ("landing.html", "text/html"), "/app": ("index.html", "text/html"),
                   "/admin": ("admin.html", "text/html")})
else:
    STATIC.update({"/classic": ("index_classic.html", "text/html"), "/index_classic.html": ("index_classic.html", "text/html")})
VIEWER_POST = {"/api/chat", "/api/chat/start", "/api/insights"}  # seules actions permises au rôle « lecteur »
CSP = ("default-src 'self'; script-src 'self' 'unsafe-inline' https://cdn.jsdelivr.net; "
       "style-src 'self' 'unsafe-inline' https://fonts.googleapis.com; font-src https://fonts.gstatic.com; "
       "img-src 'self' data: blob:; connect-src 'self' " + (os.environ.get("SUPABASE_URL", "") or "") +
       "; frame-ancestors 'none'; base-uri 'self'; form-action 'self'")

_rate, _rate_lock = {}, threading.Lock()


def _rate_ok(key, limit=30, window=60):
    now = time.time()
    with _rate_lock:
        hits = [t for t in _rate.get(key, []) if now - t < window]
        ok = len(hits) < limit
        if ok:
            hits.append(now)
        _rate[key] = hits
        return ok


def download_url(fn):
    """Lien de téléchargement. En ligne, il est signé et valable 2 heures pour l'organisation courante."""
    from urllib.parse import quote
    if not T.CLOUD:
        return f"/download/{quote(fn)}"
    org, exp = T.current().org_id, int(time.time()) + 7200
    sig = hmac.new(SECRET, f"{org}|{fn}|{exp}".encode(), hashlib.sha256).hexdigest()[:40]
    return f"/download/{quote(fn)}?k={org}.{exp}.{sig}"


def _list_exports():
    d, files = T.exports_dir(), []
    for fn in os.listdir(d):
        fp = os.path.join(d, fn)
        if os.path.isfile(fp):
            st = os.stat(fp)
            files.append({"file": fn, "url": download_url(fn), "size_kb": round(st.st_size / 1024, 1),
                          "modified": time.strftime("%Y-%m-%dT%H:%M:%S", time.localtime(st.st_mtime))})
    return sorted(files, key=lambda x: x["modified"], reverse=True)


def _safe_name(name):
    base = os.path.basename(str(name or "source"))
    stem, ext = os.path.splitext(base)
    stem = re.sub(r"[^\w\-. ]+", "_", stem, flags=re.UNICODE).strip("._ ")[:80] or "source"
    return stem + ext.lower()


class Handler(http.server.BaseHTTPRequestHandler):
    server_version = "KLDataViz"
    sys_version = ""
    timeout = 120

    # ---- réponses
    def _security_headers(self):
        self.send_header("X-Content-Type-Options", "nosniff")
        self.send_header("Referrer-Policy", "same-origin")
        if T.CLOUD:
            self.send_header("Strict-Transport-Security", "max-age=31536000; includeSubDomains")
            self.send_header("X-Frame-Options", "DENY")
            self.send_header("Content-Security-Policy", CSP)

    def _send(self, code, obj, ctype="application/json"):
        if isinstance(obj, (dict, list)):
            obj = json.dumps(_json_safe(obj), ensure_ascii=False, default=str)
        data = obj.encode("utf-8") if isinstance(obj, str) else obj
        self.send_response(code)
        self.send_header("Content-Type", ctype + "; charset=utf-8")
        if ctype == "application/json":
            self.send_header("Cache-Control", "no-store")
        elif ctype in ("text/html", "application/javascript"):  # toujours la dernière version après une mise à jour
            self.send_header("Cache-Control", "no-cache")
        self._security_headers()
        self.send_header("Content-Length", str(len(data)))
        self.end_headers()
        self.wfile.write(data)

    def _payload(self, limit=MAX_JSON):
        ln = int(self.headers.get("Content-Length") or 0)
        if ln > limit:
            raise ValueError(f"Envoi trop volumineux (maximum {limit // (1024 * 1024)} Mo).")
        raw = self.rfile.read(ln) if ln else b""
        try:
            return json.loads(raw.decode("utf-8") or "{}")
        except Exception:
            return {}

    def do_GET(self):
        self._dispatch("GET")

    def do_POST(self):
        self._dispatch("POST")

    def log_message(self, *a):
        pass

    # ---- aiguillage : fichiers publics, puis connexion, puis organisation
    def _local_origin_ok(self):
        """Version locale (sans comptes) : seules les pages servies par la console elle-même peuvent l'appeler.
        Bloque un site web malveillant ouvert dans le navigateur (requête vers 127.0.0.1) et le « DNS rebinding »."""
        local = ("127.0.0.1", "localhost", "::1")
        host = urlparse("//" + (self.headers.get("Host") or "")).hostname
        if host not in local:
            return False
        origin = self.headers.get("Origin")
        return not origin or urlparse(origin).hostname in local

    def _dispatch(self, method):
        path = urlparse(self.path).path
        try:
            if not T.CLOUD and (path.startswith(("/api/", "/download/"))) and not self._local_origin_ok():
                return self._send(403, {"error": "Requête refusée : la console locale n'accepte que ses propres pages."})
            if method == "GET" and path in STATIC:
                fn, ctype = STATIC[path]
                with open(os.path.join(BASE, fn), "rb") as f:
                    return self._send(200, f.read(), ctype)
            if method == "GET" and path == "/healthz":
                return self._send(200, {"ok": True})
            if method == "GET" and path == "/config.js":
                cfg = {"cloud": T.CLOUD}
                if T.CLOUD:
                    cfg.update(supabaseUrl=cloud.SUPABASE_URL, supabaseAnonKey=cloud.SUPABASE_ANON_KEY)
                return self._send(200, "window.KLDV=" + json.dumps(cfg) + ";", "application/javascript")
            if method == "GET" and path.startswith("/download/"):
                return self._download()
            if method == "GET" and path.startswith("/t/"):
                return self._shared(path[3:].strip("/"))
            if not path.startswith("/api/"):
                return self._send(404, {"error": "Page introuvable."})
            if T.CLOUD and method == "GET" and path.startswith("/api/public/"):
                return self._public(path)
            if not T.CLOUD:
                with T.use(T.LOCAL):
                    return self._route(method, path, self._payload(MAX_UPLOAD) if method == "POST" else {})
            auth = self.headers.get("Authorization") or ""
            user = cloud.verify_token(auth[7:] if auth.startswith("Bearer ") else "")
            payload = self._payload(MAX_UPLOAD if path == "/api/upload-source" else MAX_JSON) if method == "POST" else {}
            if path.startswith("/api/admin/"):  # administration de la plateforme : pas besoin d'espace
                actx = T.Ctx(key=user["id"], user_id=user["id"], email=user["email"],
                             is_admin=cloud.is_platform_admin(user["email"]))
                return self._admin(method, path, payload, actx)
            ctx = cloud.resolve(user)
            if ctx is None:  # compte sans espace : création de l'organisation
                if path == "/api/onboarding" and method == "POST":
                    cloud.create_org(user, payload.get("org"))
                    return self._send(200, {"ok": True})
                if path == "/api/me":
                    return self._send(200, {"onboarding": True, "email": user["email"]})
                return self._send(409, {"error": "Créez d'abord votre espace.", "onboarding": True})
            if method == "POST" and ctx.role == "viewer" and path not in VIEWER_POST:
                raise T.Forbidden("Votre rôle (lecteur) permet de consulter et de poser des questions uniquement.")
            with T.use(ctx):
                return self._route(method, path, payload)
        except T.AuthError as e:
            self._send(401, {"error": str(e), "auth": True})
        except T.Forbidden as e:
            self._send(403, {"error": str(e)})
        except T.QuotaExceeded as e:
            self._send(402, {"error": str(e), "quota": True})
        except T.Unavailable as e:
            self._send(503, {"error": str(e), "unavailable": True})
        except PermissionError:
            self._send(403, {"error": "Accès refusé."})
        except (ValueError, KeyError, FileNotFoundError, netguard.BlockedURL) as e:
            self._send(400, {"error": str(e).strip("'")})
        except Exception as e:
            # En ligne, le détail reste dans les journaux ; la référence permet de l'y retrouver.
            ref = secrets.token_hex(3)
            sys.stderr.write(f"[erreur {ref}] {method} {path} : {type(e).__name__} : {e}\n"
                             + "".join(traceback.format_exception(e)[-4:]))
            self._send(500, {"error": f"Erreur interne (réf. {ref}). Réessayez." if T.CLOUD else f"{type(e).__name__} : {e}"})

    def _shared(self, token):
        """Tableau de bord partagé par lien (sans compte) : dernière version enregistrée, page isolée."""
        hit = shares.resolve(token)
        html = None
        if hit:
            org, sid = hit
            ctx = T.LOCAL if org == "local" else T.Ctx(key="share", org_id=org)
            with T.use(ctx):
                try:
                    html = sources.get(sid).get("html")
                except KeyError:
                    html = None
        if not html:
            body = ("<!doctype html><meta charset=utf-8><title>Lien indisponible</title>"
                    "<body style='font:16px system-ui;padding:48px;color:#11241A'><h1>Lien indisponible</h1>"
                    "<p>Ce lien de partage a expiré ou a été révoqué. Demandez un nouveau lien à la personne qui vous l'a envoyé.</p>")
            data, code = body.encode("utf-8"), 404
        else:
            data, code = html.encode("utf-8"), 200
        self.send_response(code)
        self.send_header("Content-Type", "text/html; charset=utf-8")
        self.send_header("Cache-Control", "no-store")
        self.send_header("X-Robots-Tag", "noindex, nofollow")
        self.send_header("Referrer-Policy", "no-referrer")
        self.send_header("X-Content-Type-Options", "nosniff")
        # Page isolée : origine opaque (aucun accès aux cookies ni au stockage de l'application)
        self.send_header("Content-Security-Policy",
                         "sandbox allow-scripts allow-popups; default-src 'none'; script-src 'unsafe-inline' https://cdn.jsdelivr.net; "
                         "style-src 'unsafe-inline' https://fonts.googleapis.com; font-src https://fonts.gstatic.com; img-src data:; "
                         "frame-ancestors 'none'")
        self.send_header("Content-Length", str(len(data)))
        self.end_headers()
        self.wfile.write(data)

    def _download(self):
        u = urlparse(self.path)
        fn = os.path.basename(urllib.request.url2pathname(u.path))
        qs = parse_qs(u.query)
        inline = qs.get("inline") == ["1"]
        if T.CLOUD:
            try:
                org, exp, sig = (qs.get("k") or [""])[0].split(".")
                good = hmac.new(SECRET, f"{org}|{fn}|{exp}".encode(), hashlib.sha256).hexdigest()[:40]
                if not hmac.compare_digest(sig, good) or int(exp) < time.time():
                    raise ValueError
            except ValueError:
                return self._send(403, {"error": "Lien expiré. Rouvrez la page Rapports."})
            ctx = T.Ctx(key="download", org_id=org)
        else:
            ctx = T.LOCAL
        with T.use(ctx):
            fp = os.path.join(T.exports_dir(), fn)
        if not os.path.isfile(fp):
            return self._send(404, {"error": "fichier introuvable"})
        ctype = {".pdf": "application/pdf", ".html": "text/html",
                 ".xlsx": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"}.get(
            os.path.splitext(fn)[1], "application/octet-stream")
        with open(fp, "rb") as f:
            data = f.read()
        self.send_response(200)
        self.send_header("Content-Type", ctype)
        self.send_header("Content-Disposition", f'{"inline" if inline else "attachment"}; filename="{fn}"')
        self._security_headers()
        if T.CLOUD and ctype == "text/html":  # rapport HTML ouvert seul : aucun script exécuté
            self.send_header("Content-Security-Policy", "default-src 'none'; style-src 'unsafe-inline'; img-src data:")
        self.send_header("Content-Length", str(len(data)))
        self.end_headers()
        self.wfile.write(data)

    # ---- routes de l'API (contexte de l'organisation déjà posé)
    def _route(self, method, path, p):
        ctx, st = T.current(), T.state()
        if method == "GET":
            if path == "/api/me":
                return self._send(200, cloud.me(ctx) if T.CLOUD else {"local": True})
            if path == "/api/status":
                if T.CLOUD:
                    return self._send(200, {"status": "CLOUD"})
                code, r = orchestrator("/api/v1/agents/me")
                return self._send(code if code else 502, r)
            if path == "/api/tasks":
                if T.CLOUD:
                    return self._send(200, [])
                return self._send(200, _hermes_tasks())
            if path == "/api/exports":
                return self._send(200, _list_exports())
            if path in ("/api/architect/job", "/api/chat/job", "/api/hermes/job"):
                return self._send(200, job_status((parse_qs(urlparse(self.path).query).get("id") or [""])[0]))
            if path == "/api/settings":
                return self._send(200, studio.public_settings())
            if path == "/api/schedules":
                return self._send(200, studio.list_schedules())
            if path == "/api/sources":
                return self._send(200, sources.list_sources())
            if path == "/api/dictionnaire":
                return self._send(200, dictionnaire.info())
            if path == "/api/shares":
                return self._send(200, shares.list_for((parse_qs(urlparse(self.path).query).get("id") or [""])[0]))
            if path == "/api/sources/get":
                return self._send(200, sources.get((parse_qs(urlparse(self.path).query).get("id") or [""])[0]))
            if path == "/api/ingest/versions":
                ds_id = (parse_qs(urlparse(self.path).query).get("id") or [""])[0]
                reg = T.registry()
                reg.ensure()
                import registry_meta as RM
                vm = RM.load(_registry_meta_dir(), ds_id).get("versions", {})
                return self._send(200, [dict(v, changements=(vm.get(str(v.get("version_no"))) or {}).get("changements"))
                                        for v in reg.list_versions(ds_id)])
            if path == "/api/fetch-source":
                if T.CLOUD:
                    return self._send(200, {"source": None, "message": "Indisponible dans la version en ligne."})
                return self._send(200, fetch_source())
            if path == "/api/ingest/list":
                if not INGESTION_OK:
                    return self._send(500, {"error": f"Agent d'ingestion indisponible : {INGESTION_ERROR}"})
                reg = T.registry()
                info = reg.ensure()
                import registry_meta as RM
                md = _registry_meta_dir()
                datasets = [dict(x, meta=RM.summary(RM.load(md, x.get("dataset_id")))) for x in reg.latest_versions()]
                return self._send(200, {"backend": reg.backend, "connection": {} if T.CLOUD else info,
                                        "summary": reg.summary(), "datasets": datasets})
            if T.CLOUD and path == "/api/team":
                return self._send(200, cloud.team(ctx))
            if T.CLOUD and path == "/api/announcements":
                return self._send(200, cloud.announcements("app"))
            return self._send(404, {"error": "not found"})

        # ---- POST
        if path in ("/api/analyze", "/api/chat", "/api/chat/start", "/api/insights", "/api/clean", "/api/architect",
                    "/api/architect/apply", "/api/sheets/understand", "/api/hermes/charts", "/api/hermes/clean",
                    "/api/dashboard") and not _rate_ok(f"{ctx.key}:{path}"):
            return self._send(429, {"error": "Trop de demandes en peu de temps. Patientez une minute."})
        if path == "/api/analyze":
            file_path = str(p.get("file_path") or "")
            if T.CLOUD and not os.path.realpath(file_path).startswith(os.path.realpath(T.cleaned_dir()) + os.sep):
                cloud.consume(ctx, "analyses")  # la réanalyse d'une copie corrigée n'est pas décomptée
            result = analyze(file_path)
            if "error" not in result:
                st["last_result"], st["current_path"], st["result_path"] = result, file_path, file_path
            return self._send(200, result)
        if path == "/api/dashboard":
            if T.CLOUD and _llm_config():
                cloud.consume(ctx, "ai")
            return self._send(200, do_dashboard(str(p.get("file_path") or ""), p.get("name")))
        if path in ("/api/dictionnaire/import", "/api/dictionnaire/delete"):
            if T.CLOUD:
                cloud.require(T.current(), "owner", "analyst")
            if path.endswith("/delete"):
                return self._send(200, dictionnaire.supprimer())
            raw = str(p.get("data") or "")
            if len(raw) > 1024 * 1024 * 4 // 3 + 4:  # taille vérifiée avant décodage
                raise ValueError("Dictionnaire trop volumineux (1 Mo maximum).")
            data = base64.b64decode(raw)
            if not data:
                raise ValueError("Fichier vide.")
            return self._send(200, dictionnaire.importer(data, str(p.get("name") or "")))
        if path == "/api/dhis2/connect":  # identifiants chiffrés dans l'espace, puis lecture d'essai
            return self._send(200, dhis2_connect(p))
        if path == "/api/sources/refresh":  # relit la source, mêmes indicateurs : pas de quota IA
            return self._send(200, sources.refresh(str(p.get("id")), _read_df, build_dashboard))
        if path == "/api/sources/freq":
            return self._send(200, sources.set_freq(str(p.get("id")), str(p.get("freq"))))
        if path == "/api/sources/delete":
            sources.delete(str(p.get("id")))
            shares.revoke_all(str(p.get("id")))
            return self._send(200, {"ok": True})
        if path == "/api/shares/create":
            sources.get(str(p.get("id")), with_html=False)  # le tableau doit exister dans cet espace
            return self._send(200, shares.create(str(p.get("id")), int(p.get("days") or 30)))
        if path == "/api/shares/revoke":
            shares.revoke(str(p.get("token")))
            return self._send(200, {"ok": True})
        if path == "/api/sources/mail":
            return self._send(200, sources.set_mail(str(p.get("id")), p.get("to"), str(p.get("freq") or "off")))
        if path == "/api/sources/send":
            return self._send(200, sources.send_mail_now(str(p.get("id")), _read_df, build_dashboard, p.get("to") or None))
        if path == "/api/clean":
            hermes = p.get("hermes") if isinstance(p.get("hermes"), list) else None
            return self._send(200, clean(str(p.get("file_path") or ""), p.get("skip") or None, hermes))
        if path == "/api/registry/describe":
            import registry_meta as RM
            df, err = _read_df(str(p.get("file_path") or ""))
            if err:
                return self._send(200, {"error": err})
            meta = RM.describe(_registry_meta_dir(), p.get("dataset_id"), p.get("version_no"), df, _period_series,
                               p.get("comprehension") if isinstance(p.get("comprehension"), dict) else None,
                               str(p.get("name") or "")[:200])
            return self._send(200, {"meta": RM.summary(meta)})
        if path == "/api/hermes/clean":
            if T.CLOUD and _llm_config():
                cloud.consume(ctx, "ai")
            fp = str(p.get("file_path") or "")
            return self._send(200, start_job(lambda progress: hermes_clean(fp, progress)))
        if path == "/api/hermes/charts":
            if T.CLOUD and _llm_config():
                cloud.consume(ctx, "ai")
            fp = str(p.get("file_path") or "")
            comp = p.get("comprehension")
            return self._send(200, start_job(lambda progress: hermes_charts(fp, progress, comp)))
        if path == "/api/architect":
            if T.CLOUD and _llm_config():
                cloud.consume(ctx, "ai")
            return self._send(200, architect_start(_architect_paths(p)))
        if path == "/api/architect/apply":
            return self._send(200, architect_apply(_architect_paths(p), p.get("spec")))
        if path == "/api/sheets":
            return self._send(200, sheets_info(str(p.get("file_path") or "")))
        if path == "/api/sheets/understand":
            if T.CLOUD and _llm_config():
                cloud.consume(ctx, "ai")
            fp = str(p.get("file_path") or "")
            names = p.get("sheets") if isinstance(p.get("sheets"), list) else []
            return self._send(200, start_job(lambda progress: workbook_understand(fp, names, progress)))
        if path == "/api/sheets/extract":
            sheets = p.get("sheets") if isinstance(p.get("sheets"), list) else []
            return self._send(200, sheets_extract(str(p.get("file_path") or ""), sheets, bool(p.get("empiler"))))
        if path == "/api/ingest":
            return self._ingest(p)
        if path in ("/api/export/pdf", "/api/export/excel", "/api/export/html"):
            if st["last_result"] is None:
                return self._send(400, {"error": "Lancez d'abord une analyse."})
            result = st["last_result"]
            if path.endswith("/pdf") and isinstance(p.get("charts"), list):
                # Graphiques dessinés par l'interface (identiques à l'écran Analyser), 12 au plus
                result = dict(result, _chart_images=[c for c in p["charts"][:12]
                                                     if isinstance(c, dict) and isinstance(c.get("image"), str)])
            if path.endswith("/pdf") and p.get("hermes_synthese"):
                result = dict(result, _hermes_synthese=str(p["hermes_synthese"])[:900])
            fn = EXPORTERS[path.rsplit("/", 1)[-1]](result)
            return self._send(200, {"url": download_url(fn), "file": fn})
        if path == "/api/exemple":  # « Essayer avec un exemple » : jeu FICTIF généré à la volée, rien dans le dépôt
            import exemple
            return self._send(200, exemple.ecrire(T.uploads_dir()))
        if path == "/api/upload-source":
            name = _safe_name(p.get("name"))
            if T.CLOUD and not name.lower().endswith(ALLOWED_EXT):
                raise ValueError("Format non pris en charge. Utilisez CSV, Excel, JSON ou Parquet.")
            data = base64.b64decode(p.get("data") or "")
            if len(data) > MAX_UPLOAD:
                raise ValueError(f"Fichier trop volumineux (maximum {MAX_UPLOAD // (1024 * 1024)} Mo).")
            dest = os.path.join(T.uploads_dir(), name)
            with open(dest, "wb") as f:
                f.write(data)
            return self._send(200, {"path": dest, "name": name, "size": len(data)})
        if path in ("/api/chat", "/api/chat/start"):
            question = str(p.get("question") or "")[:2000]
            # Fichier en cours envoyé par le navigateur : le chat ne dépend plus de la mémoire du
            # serveur, effacée à chaque redémarrage/déploiement. _read_df vérifie qu'il appartient à l'organisation.
            if p.get("file_path"):
                st["current_path"] = str(p["file_path"])
            if T.CLOUD and _llm_config():
                cloud.consume(ctx, "ai")
            history = p.get("historique") if isinstance(p.get("historique"), list) else None

            def answer(progress=None):
                linked = p.get("fichiers_lies") if isinstance(p.get("fichiers_lies"), list) else None
                resp = llm_chat(question, p.get("contexte"), history, progress, linked)
                return resp if resp is not None else chat_query(question)  # repli par règles sans IA
            if path == "/api/chat/start":  # plusieurs allers-retours avec l'IA : tâche de fond suivie en direct
                return self._send(200, start_job(answer))
            return self._send(200, answer())
        if path == "/api/insights":
            if st["last_result"] is not None and T.CLOUD and _llm_config():
                cloud.consume(ctx, "ai")
            ins = studio.insights(st["last_result"], _llm_call)
            if ins and st["last_result"] is not None:
                st["last_result"]["_insights"] = ins  # repris dans le rapport PDF ; remplacé à la prochaine analyse
            return self._send(200, {"insights": ins})
        if path == "/api/settings":
            return self._send(200, studio.save_settings(p))
        if path == "/api/schedules":
            return self._send(200, studio.create_schedule(p))
        if path == "/api/schedules/update":
            return self._send(200, studio.update_schedule(p.get("id"), on=p.get("on")))
        if path == "/api/schedules/delete":
            studio.delete_schedule(p.get("id"))
            return self._send(200, {"ok": True})
        if path == "/api/schedules/run":
            if T.CLOUD:
                cloud.consume(ctx, "analyses")
            return self._send(200, studio.run_schedule(p.get("id"), analyze, EXPORTERS, T.exports_dir(), schedule_commentary))
        if path == "/api/exports/delete":
            fp = os.path.join(T.exports_dir(), os.path.basename(str(p.get("file") or "")))
            if not os.path.isfile(fp):
                raise FileNotFoundError("Fichier introuvable.")
            os.remove(fp)
            return self._send(200, {"ok": True})
        if T.CLOUD and path == "/api/onboarding":
            raise ValueError("Votre espace existe déjà.")
        if T.CLOUD and path == "/api/team/invite":
            return self._send(200, cloud.invite(ctx, p.get("email"), p.get("role") or "analyst"))
        if T.CLOUD and path == "/api/team/role":
            cloud.set_role(ctx, str(p.get("user_id")), p.get("role"))
            return self._send(200, {"ok": True})
        if T.CLOUD and path == "/api/team/remove":
            cloud.remove_member(ctx, str(p.get("user_id")))
            return self._send(200, {"ok": True})
        if T.CLOUD and path == "/api/team/cancel":
            cloud.cancel_invite(ctx, str(p.get("id")))
            return self._send(200, {"ok": True})
        if path == "/api/heartbeat" and not T.CLOUD:
            code, r = orchestrator("/api/v1/agents/heartbeat", method="POST", body={})
            return self._send(code if code else 502, r)
        return self._send(404, {"error": "not found"})

    def _public(self, path):
        """Données publiques de la page d'accueil (sans connexion)."""
        if path == "/api/public/plans":
            return self._send(200, cloud.public_plans())
        if path == "/api/public/site":
            return self._send(200, cloud.site_settings())
        if path == "/api/public/announcements":
            return self._send(200, cloud.announcements("public"))
        return self._send(404, {"error": "not found"})

    def _admin(self, method, path, p, ctx):
        """API de l'administration de la plateforme (réservée aux administrateurs)."""
        if not ctx.is_admin:
            raise T.Forbidden("Réservé aux administrateurs de la plateforme.")
        q = parse_qs(urlparse(self.path).query)
        arg = lambda k: (q.get(k) or [""])[0]
        if method == "GET":
            routes = {"/api/admin/me": lambda: {"email": ctx.email, "is_admin": True},
                      "/api/admin/stats": lambda: cloud.admin_stats(ctx),
                      "/api/admin/orgs": lambda: cloud.admin_orgs(ctx, arg("q")),
                      "/api/admin/org": lambda: cloud.admin_org_detail(ctx, arg("id")),
                      "/api/admin/plans": lambda: cloud.admin_plans(ctx),
                      "/api/admin/admins": lambda: cloud.admin_admins(ctx),
                      "/api/admin/announcements": lambda: cloud.admin_announcements(ctx),
                      "/api/admin/site": lambda: cloud.site_settings()}
        else:
            routes = {"/api/admin/plan": lambda: cloud.admin_set_plan(ctx, str(p.get("org_id")), p.get("plan"), p.get("trial_days")),
                      "/api/admin/extend": lambda: cloud.admin_extend_trial(ctx, str(p.get("org_id")), p.get("days") or 14),
                      "/api/admin/suspend": lambda: cloud.admin_suspend(ctx, str(p.get("org_id")), p.get("suspended")),
                      "/api/admin/invite-client": lambda: cloud.admin_invite_client(ctx, p.get("org"), p.get("email"), p.get("plan")),
                      "/api/admin/resend-invite": lambda: cloud.admin_resend_invite(ctx, str(p.get("org_id")), p.get("email")),
                      "/api/admin/test-mail": lambda: cloud.admin_test_mail(ctx),
                      "/api/admin/plans/save": lambda: cloud.admin_plan_save(ctx, p),
                      "/api/admin/plans/delete": lambda: cloud.admin_plan_delete(ctx, str(p.get("key"))),
                      "/api/admin/admins/add": lambda: cloud.admin_admin_add(ctx, p.get("email")),
                      "/api/admin/admins/remove": lambda: cloud.admin_admin_remove(ctx, p.get("email")),
                      "/api/admin/announcements/save": lambda: cloud.admin_announcement_save(ctx, p),
                      "/api/admin/announcements/delete": lambda: cloud.admin_announcement_delete(ctx, str(p.get("id"))),
                      "/api/admin/site": lambda: cloud.admin_site_save(ctx, p)}
        fn = routes.get(path)
        if fn is None:
            return self._send(404, {"error": "not found"})
        out = fn()
        return self._send(200, out if out is not None else {"ok": True})

    def _ingest(self, p):
        if not INGESTION_OK:
            return self._send(500, {"error": f"Agent d'ingestion indisponible : {INGESTION_ERROR}"})
        source = str(p.get("file_path") or p.get("source") or "").strip()
        if not source:
            raise ValueError("Indiquez un fichier ou un lien.")
        kwargs = {"registry_schema": T.current().schema}
        if T.storage_root():
            kwargs["storage_root"] = T.storage_root()
        if source.startswith("dhis2://"):  # lu par le connecteur, versionné comme un fichier CSV de l'espace
            df, err = _read_df(source)
            if err:
                raise ValueError(err)
            local = os.path.join(T.uploads_dir(), source_label(source) + ".csv")
            df.to_csv(local, index=False)
            mark_agrege(local, bool(df.attrs.get("agrege")))
            kwargs["dataset_name"] = os.path.basename(local)
            source = local
        elif source.startswith(("http://", "https://")):
            source = normalize_source_url(source)
            if T.CLOUD:  # téléchargé par le serveur (protection SSRF), puis ingéré comme fichier
                data, ctype = netguard.fetch(source)
                ext = ".xlsx" if data[:2] == b"PK" else ".json" if data[:1] in (b"{", b"[") else ".csv"
                local = os.path.join(T.uploads_dir(), source_label(source) + ext)
                with open(local, "wb") as f:
                    f.write(data)
                kwargs["dataset_name"] = source_label(source) + ext
                source = local
        elif T.CLOUD and not T.current().owns(source):
            raise PermissionError("source hors de l'espace")
        if T.CLOUD and p.get("kind") in ("postgres", "sqlite"):
            raise ValueError("L'import depuis une base de données n'est pas disponible dans la version en ligne.")
        sheet = p.get("sheet")
        if sheet not in (None, ""):
            try:
                sheet = int(sheet)
            except (TypeError, ValueError):
                pass
            kwargs["sheet"] = sheet
        for key, arg in (("kind", "kind"), ("table", "table"), ("schema", "schema_name"), ("name", "dataset_name")):
            if p.get(key):
                kwargs[arg] = p[key]
        if p.get("limit"):
            kwargs["limit"] = int(p["limit"])
        if p.get("force"):
            kwargs["force"] = True
        try:
            result = ingest_agent(source, **kwargs)
        except Exception as exc:
            result = {"error": f"{type(exc).__name__} : {exc}"}
        if "error" not in result:
            T.state()["last_ingest"] = result
        return self._send(200, result)


def heartbeat_loop():
    while True:
        try:
            orchestrator("/api/v1/agents/heartbeat", method="POST", body={})
        except Exception:
            pass
        time.sleep(45)


if __name__ == "__main__":
    if not T.CLOUD:
        threading.Thread(target=heartbeat_loop, daemon=True).start()
    studio.start_scheduler(analyze, EXPORTERS, T.LOCAL.dir("exports") if not T.CLOUD else None, schedule_commentary)
    sources.start_auto_refresh(_read_df, build_dashboard, cloud.org_ctx if T.CLOUD else None)
    host = "0.0.0.0" if T.CLOUD else "127.0.0.1"
    port = int(os.environ.get("PORT") or PORT)
    socketserver.ThreadingTCPServer.allow_reuse_address = True
    socketserver.ThreadingTCPServer.daemon_threads = True
    with socketserver.ThreadingTCPServer((host, port), Handler) as httpd:
        print(f"KLDataViz {'en ligne' if T.CLOUD else 'Console'} démarrée : http://localhost:{port}")
        httpd.serve_forever()
