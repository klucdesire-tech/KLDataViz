# -*- coding: utf-8 -*-
"""Non-régression de la console locale : python -m pytest tests/test_local.py"""
import os
import time

import pandas as pd
import pytest

os.environ.pop("KLDV_MODE", None)
from _harness import call, reset_app_modules, start  # noqa: E402
reset_app_modules()
import server  # noqa: E402


@pytest.fixture(scope="module")
def base(tmp_path_factory):
    httpd, url = start(server)
    yield url
    httpd.shutdown()


@pytest.fixture(scope="module")
def csv(tmp_path_factory):
    p = tmp_path_factory.mktemp("d") / "ventes_test.csv"
    pd.DataFrame({"region": ["A", "B", "A", "C"] * 10, "montant": range(40)}).to_csv(p, index=False)
    return str(p)


def test_pages_publiques(base):
    assert call(base, "/", raw=True)[0] == 200
    assert call(base, "/classic", raw=True)[0] == 200
    assert call(base, "/healthz")[1] == {"ok": True}


def test_parcours_analyse_export(base, csv):
    code, a = call(base, "/api/analyze", {"file_path": csv})
    assert code == 200 and a["rows"] == 40, a
    code, c = call(base, "/api/clean", {"file_path": csv})
    assert code == 200 and "cleaned_path" in c
    code, e = call(base, "/api/export/html", {})
    assert code == 200 and e["url"].startswith("/download/")
    code, data = call(base, e["url"], raw=True)
    assert code == 200 and len(data) > 100
    code, files = call(base, "/api/exports")
    assert any(f["file"] == e["file"] for f in files)
    assert call(base, "/api/exports/delete", {"file": e["file"]})[0] == 200
    os.remove(c["cleaned_path"])


def test_architecte_format_large(base, tmp_path, monkeypatch):
    """Export type DHIS2 (« ; », dimensions dans les en-têtes, ligne d'agrégat) → format long."""
    monkeypatch.setattr(server, "_llm_config", lambda: None)  # repli par règles, sans appel réseau
    ages = ["00-04 ans", "05-09 ans", "50 ans et plus"]
    cols = {f"Mars 2026 SIG - Nombre de {i} {a}": [3, 5, 7, 3] for i in ("consultations", "consultants") for a in ages}
    p = tmp_path / "dhis2.csv"
    pd.DataFrame({"orgid": ["a1", "b2", "c3", "d4"], "nom": ["CSR-D A", "CSR-D B", "CSU C", "DISTRICT"], **cols}) \
        .to_csv(p, sep=";", index=False, encoding="utf-8-sig")
    code, job = call(base, "/api/architect", {"file_path": str(p)})
    assert code == 200 and job.get("job"), job
    for _ in range(100):  # tâche de fond
        code, j = call(base, f"/api/architect/job?id={job['job']}")
        if j.get("status") == "done":
            break
        time.sleep(0.1)
    r = j["result"]
    assert not r.get("error"), r
    assert call(base, "/api/architect/job?id=inconnu")[1].get("error")
    assert r["ia"] is False and set(r["proposition"]["comprehension"]["dimensions"]) == {"periode", "tranche_age"}
    assert r["apercu"]["fait"]["lignes"] == 4 * len(cols)
    code, a = call(base, "/api/architect/apply", {"file_path": str(p), "spec": r["proposition"]["spec"]})
    assert code == 200 and a["url"].startswith("/download/"), a
    # La structure est analysable directement : vue à plat de la table de faits
    vue = a["vues"][0]
    assert vue["lignes"] == r["apercu"]["fait"]["lignes"]
    code, an = call(base, "/api/analyze", {"file_path": vue["path"]})
    assert code == 200 and an["rows"] == vue["lignes"], an
    os.remove(vue["path"])
    assert call(base, "/api/exports/delete", {"file": a["file"]})[0] == 200


def test_detail_des_corrections(base, tmp_path):
    """Chaque correction proposée est détaillée : lignes, avant → après, valeurs non convertibles."""
    p = tmp_path / "sale.csv"
    pd.DataFrame({"client": ["A", "B", "A"] + list("CDEFGHIJKLMNOPQRSTUVW"),
                  "montant": ["1 200,5", "30", "1 200,5", "inconnu"] + [str(i) for i in range(20)]}).to_csv(p, index=False)
    code, c = call(base, "/api/clean", {"file_path": str(p)})
    assert code == 200, c
    d = c["rapport"]["details"]
    assert d["doublons"]["groupes"][0]["lignes"] == [2, 4]  # n° de ligne du fichier, en-tête = 1
    assert d["type:montant"]["non_convertibles"] == 1
    assert d["type:montant"]["exemples_non_convertibles"][0] == {"ligne": 5, "valeur": "inconnu"}
    assert d["imput:montant"]["lignes"] == [5]
    os.remove(c["cleaned_path"])


def test_architecte_fichier_long_sans_ia(base, tmp_path, monkeypatch):
    """Fichier déjà « long » et IA indisponible : une structure exécutable est quand même proposée."""
    monkeypatch.setattr(server, "_llm_config", lambda: None)
    p = tmp_path / "prelevements.csv"
    pd.DataFrame({"patient": [f"P{i}" for i in range(30)], "site": ["Buyo", "Soubré", "Méagui"] * 10,
                  "resultat": range(30), "vide": [None] * 30}).to_csv(p, index=False)
    r = _wait_job(base, call(base, "/api/architect", {"file_path": str(p)})[1])
    assert r["source_structure"] == "regles" and r["proposition"]["spec"], r
    assert r["apercu"]["fait"]["colonnes"] == ["patient", "site", "resultat"]  # colonne vide retirée


def test_architecte_ia_corrige_sa_recette():
    """Recette de l'IA inexécutable → l'erreur lui est renvoyée, la recette corrigée est retenue."""
    import json as _json
    import architect
    df = pd.DataFrame({"site": ["A", "B"], "valeur": [1, 2]})
    good = {"comprehension": {"resume": "ok"}, "problemes": [], "structure": [],
            "spec": {"fact_name": "fait", "drop_columns": ["colonne_inexistante"],
                     "dimensions": {"dim_site": {"key": "SITE"}}}}
    answers = iter([_json.dumps(good), _json.dumps({"spec": {"fact_name": "fait",
                                                              "dimensions": {"dim_site": {"key": "site"}}}})])
    r = architect.propose(df, {}, lambda messages, max_tokens: next(answers))
    assert r["source_structure"] == "ia_corrigee" and "SITE" in r["spec_error"], r
    assert set(r["apercu"]) == {"fait", "dim_site"}


def test_csv_point_virgule(base, tmp_path):
    """Exports français / DHIS2 : séparateur « ; » et encodage Windows détectés à l'analyse."""
    p = tmp_path / "export_fr.csv"
    pd.DataFrame({"établissement": ["CSR Buyo", "CSU Méagui"] * 10, "consultations": range(20)}) \
        .to_csv(p, sep=";", index=False, encoding="cp1252")
    code, a = call(base, "/api/analyze", {"file_path": str(p)})
    assert code == 200 and a["columns"] == 2, a


def test_reparation_encodage():
    """« contrÃ´le » → « contrôle », détaillé, et refusable par l'utilisateur."""
    import modeling
    df = pd.DataFrame({"motif": ["sous contrÃ´le ARV", "DÃ©cÃ¨s", "Soubré", None] * 5})
    out, r = modeling.clean_dataframe(df)
    assert r["encodage_corrige"] == {"motif": 2}  # après dédoublonnage des 5 répétitions
    assert out["motif"].iloc[0] == "sous contrôle ARV" and out["motif"].iloc[2] == "Soubré"
    assert r["details"]["encodage:motif"]["exemples"][0] == {"ligne": 2, "avant": "sous contrÃ´le ARV",
                                                            "apres": "sous contrôle ARV"}
    out2, r2 = modeling.clean_dataframe(df, skip=["encodage:motif"])
    assert r2["encodage_corrige"] == {} and out2["motif"].iloc[0] == "sous contrÃ´le ARV"


def test_pas_de_fausse_correction_trim():
    """Une colonne texte avec des cases vides mais sans espaces superflus n'est pas « à corriger »."""
    import modeling
    assert modeling._coerce_series(pd.Series(["a", None, "b"]))[1] is None
    out, t = modeling._coerce_series(pd.Series([" a", None, "b "]))
    assert t == "texte (trim)" and out.iloc[0] == "a" and pd.isna(out.iloc[1])


def _wait_job(base, job):
    for _ in range(100):
        j = call(base, f"/api/architect/job?id={job['job']}")[1]
        if j.get("status") == "done":
            return j["result"]
        time.sleep(0.1)
    raise AssertionError("tâche non terminée")


def test_architecte_plusieurs_fichiers(base, tmp_path, monkeypatch):
    """Deux fichiers reliés par un identifiant de nom différent → relation détectée, modèle commun."""
    monkeypatch.setattr(server, "_llm_config", lambda: None)
    ventes = tmp_path / "ventes.csv"
    clients = tmp_path / "clients.xlsx"
    pd.DataFrame({"id_client": ["C1", "C2", "C3", "C1"] * 5, "montant": range(20)}).to_csv(ventes, index=False)
    pd.DataFrame({"CODE CLIENT": ["C1", "C2", "C3", "C4"], "ville": ["Abidjan", "Soubré", "Méagui", "Buyo"]}) \
        .to_excel(clients, index=False)
    paths = [str(ventes), str(clients)]
    r = _wait_job(base, call(base, "/api/architect", {"file_paths": paths})[1])
    assert not r.get("error"), r
    rel = r["relations_mesurees"][0]
    assert {rel["gauche"], rel["droite"]} == {"ventes.csv :: id_client", "clients.xlsx :: CODE CLIENT"}
    assert rel["cardinalite"] in ("N:1", "1:N")
    a = call(base, "/api/architect/apply", {"file_paths": paths, "spec": r["proposition"]["spec"]})[1]
    assert a.get("url") and len(a["tables"]) == 2, a
    assert call(base, "/api/exports/delete", {"file": a["file"]})[0] == 200


def test_chat_seconde_chance_code_ia(base, tmp_path, monkeypatch):
    """Code de l'IA invalide (apostrophe dans un nom de colonne) → erreur renvoyée à l'IA, code corrigé exécuté."""
    import json as _json
    p = tmp_path / "utilisateurs.csv"
    pd.DataFrame({"Nom d'utilisateur": ["awa", "koffi", "awa"], "Age (ans)": [30, 40, 50]}).to_csv(p, index=False)
    assert call(base, "/api/analyze", {"file_path": str(p)})[0] == 200
    broken = "df = data['utilisateurs']\nresult = df.groupby('Nom d'utilisateur')['Age (ans)'].mean()"
    fixed = "df = data['utilisateurs']\nresult = df.groupby(\"Nom d'utilisateur\")[\"Age (ans)\"].mean()"
    answers = iter([_json.dumps({"action": "calcul", "pourquoi": "âge moyen", "code": broken}),
                    _json.dumps({"action": "calcul", "pourquoi": "âge moyen (corrigé)", "code": fixed}),
                    _json.dumps({"action": "reponse", "answer": "Âge moyen : 40 ans pour awa et koffi."})])
    seen = []
    monkeypatch.setattr(server, "_llm_config", lambda: {"api_key": "test"})
    monkeypatch.setattr(server, "_llm_call", lambda messages, *a, **k: (seen.append(list(messages)), next(answers))[1])
    code, r = call(base, "/api/chat", {"question": "âge moyen par Nom d'utilisateur", "file_path": str(p)})
    assert code == 200 and r["type"] == "table", r
    assert r["table"]["rows"] == [["awa", "40.0"], ["koffi", "40.0"]]
    assert "unterminated string" in seen[1][-1]["content"]  # l'erreur a bien été transmise à l'IA
    assert [e["ok"] for e in r["etapes"]] == [False, True]


def test_chat_agent_plusieurs_etapes_et_historique(base, tmp_path, monkeypatch):
    """L'agent lit le résultat réel de chaque calcul avant de répondre ; l'historique et la progression circulent."""
    import json as _json
    p = tmp_path / "cv.csv"
    pd.DataFrame({"district": ["Buyo", "Soubré"] * 20, "cv": [100, 5000] * 20}).to_csv(p, index=False)
    answers = iter([
        _json.dumps({"action": "calcul", "pourquoi": "moyenne par district",
                     "code": "df = data['cv']\nresult = df.groupby(\"district\")[\"cv\"].mean()\n"
                             "chart = {'type':'bar','title':'CV','labels':list(result.index),'values':[result.tolist()]}"}),
        _json.dumps({"action": "calcul", "pourquoi": "test de Student",
                     "code": "df = data['cv']\nr = stats.ttest_ind(df[df.district==\"Buyo\"].cv, df[df.district==\"Soubré\"].cv)\n"
                             "print('p', r.pvalue)"}),
        _json.dumps({"action": "reponse", "answer": "Soubré est nettement plus élevé (test de Student)."}),
    ])
    seen = []
    monkeypatch.setattr(server, "_llm_config", lambda: {"api_key": "test"})
    monkeypatch.setattr(server, "_llm_call", lambda messages, *a, **k: (seen.append(list(messages)), next(answers))[1])
    job = call(base, "/api/chat/start", {"question": "Et entre districts ?", "file_path": str(p),
                                         "historique": [{"role": "user", "text": "Parle-moi de la charge virale"},
                                                        {"role": "assistant", "text": "Elle varie selon le district."}]})[1]
    r = _wait_job(base, job)
    assert r["answer"].startswith("Soubré") and len(r["etapes"]) == 2 and all(e["ok"] for e in r["etapes"]), r
    assert r["charts"][0]["labels"] == ["Buyo", "Soubré"]
    assert "Buyo | 100.0" in seen[1][-1]["content"]          # le résultat réel du calcul 1 est relu
    assert "Sortie de print()" in seen[2][-1]["content"]      # la sortie du test statistique aussi
    assert seen[0][1]["content"] == "Parle-moi de la charge virale"  # mémoire de la conversation


def test_chat_apres_redemarrage(base, tmp_path, monkeypatch):
    """Mémoire du serveur effacée (déploiement) : le fichier envoyé par le navigateur est quand même utilisé."""
    import json as _json
    p = tmp_path / "apres_redemarrage.csv"
    pd.DataFrame({"district": ["Buyo", "Méagui"], "cas": [3, 5]}).to_csv(p, index=False)
    server.T._states.clear()  # comme après un redémarrage
    prompts = []
    monkeypatch.setattr(server, "_llm_config", lambda: {"api_key": "test"})
    monkeypatch.setattr(server, "_llm_call", lambda messages, *a, **k: (prompts.append(messages[0]["content"]),
                                                                        _json.dumps({"answer": "ok"}))[1])
    code, r = call(base, "/api/chat", {"question": "résume", "file_path": str(p)})
    assert code == 200 and r["answer"] == "ok", r
    assert "apres_redemarrage" in prompts[0] and "fichier_en_cours_d_analyse" in prompts[0]


def test_dates_converties_rapidement():
    import modeling
    s = pd.Series(["01/02/2025", " 15/03/2025", None, "2025-04-30"] * 500)
    out, t = modeling._coerce_series(s)
    assert t == "date" and out.iloc[0] == pd.Timestamp("2025-02-01") and out.iloc[1] == pd.Timestamp("2025-03-15")
    assert pd.isna(out.iloc[2])


def test_pas_de_tableau_croise_geant_sur_les_identifiants():
    """Colonnes identifiants (une valeur par ligne) : pas de V de Cramér (8 Go mesurés sur 33 000 lignes)."""
    import time as _t
    import modeling
    n = 20000
    df = pd.DataFrame({"code_patient": [f"P{i}" for i in range(n)], "code_echantillon": [f"E{i}" for i in range(n)],
                       "sexe": ["M", "F"] * (n // 2), "district": ["Buyo", "Soubré", "Méagui", "Nawa"] * (n // 4)})
    t0 = _t.time()
    r = modeling.bivariate_stats(df)
    assert _t.time() - t0 < 5
    assert [(c["a"], c["b"]) for c in r["cramer"]] == [("district", "sexe")]


def test_chat_interprete_l_analyse(base, tmp_path, monkeypatch):
    """Le chat reçoit le résultat de l'analyse (recalculé après redémarrage) et le contexte de l'écran."""
    import json as _json
    p = tmp_path / "analyse_chat.csv"
    pd.DataFrame({"district": ["Buyo", "Soubré"] * 30, "age_ans": list(range(58)) + [300, 400]}).to_csv(p, index=False)
    server.T._states.clear()  # comme après un redémarrage : plus de résultat en mémoire
    prompts = []
    monkeypatch.setattr(server, "_llm_config", lambda: {"api_key": "test"})
    monkeypatch.setattr(server, "_llm_call", lambda messages, *a, **k: (prompts.append(messages[0]["content"]),
                                                                        _json.dumps({"answer": "synthèse"}))[1])
    code, r = call(base, "/api/chat", {"question": "Qu'est-ce qu'on retient de cette analyse ?", "file_path": str(p),
                                       "contexte": {"corrections_appliquees": ["2 doublons retirés"]}})
    assert code == 200 and r["answer"] == "synthèse", r
    prompt = prompts[0]
    assert "RÉSULTAT DE L'ANALYSE" in prompt and "INTERPRÉTATION" in prompt
    assert "age_ans" in prompt and "anomalies" in prompt and "2 doublons retirés" in prompt


def test_chat_ignore_une_conclusion_ecrite_avant_le_resultat(base, tmp_path, monkeypatch):
    """L'IA renvoie un calcul ET une conclusion inventée dans le même tour : seul le calcul compte."""
    import json as _json
    p = tmp_path / "invente.csv"
    pd.DataFrame({"sexe": ["F", "M"] * 10, "ok": [1, 0] * 10}).to_csv(p, index=False)
    two = (_json.dumps({"action": "calcul", "pourquoi": "taux", "code": "result = data['invente'].groupby(\"sexe\")[\"ok\"].mean()"})
           + "\n" + _json.dumps({"action": "reponse", "answer": "97,0 % contre 97,5 % (inventé)"}))
    answers = iter([two, _json.dumps({"action": "reponse", "answer": "Femmes 100 %, hommes 0 %."})])
    monkeypatch.setattr(server, "_llm_config", lambda: {"api_key": "test"})
    monkeypatch.setattr(server, "_llm_call", lambda messages, *a, **k: next(answers))
    code, r = call(base, "/api/chat", {"question": "taux par sexe ?", "file_path": str(p)})
    assert r["answer"] == "Femmes 100 %, hommes 0 %." and len(r["etapes"]) == 1, r


def _classeur(path):
    """Page de garde en 1re position, données (avec titre au-dessus) en 2e, feuille vide, référentiel masqué."""
    from openpyxl import Workbook
    wb = Workbook()
    ws = wb.active
    ws.title = "Sommaire"
    for t in ("Rapport mensuel", "Source : DHIS2", "Contact : equipe data"):
        ws.append([t])
    ws = wb.create_sheet("Données")
    ws.append(["Consultations par district — 2026"])
    ws.append([])
    ws.append(["district", "mois", "consultations", "deces"])
    for i in range(120):
        ws.append([f"D{i % 6}", f"2026-{i % 12 + 1:02d}", 100 + i, i % 3])
    wb.create_sheet("Feuil3")
    ws = wb.create_sheet("Districts")
    ws.append(["district", "region"])
    for i in range(6):
        ws.append([f"D{i}", "Nord" if i < 3 else "Sud"])
    ws.sheet_state = "hidden"
    for mois in ("Janv", "Fevr"):  # un onglet par mois, même structure : à empiler
        ws = wb.create_sheet(mois)
        ws.append(["site", "cas", "date"])
        for i in range(5):
            ws.append([f"S{i}", i * 2, f"2026-0{1 if mois == 'Janv' else 2}-1{i}"])
    wb.save(path)
    return str(path)


def test_classeur_plusieurs_feuilles(base, tmp_path, monkeypatch):
    """La feuille de données est choisie (pas la page de garde), et chaque feuille peut être extraite."""
    xl = _classeur(tmp_path / "rapport.xlsx")
    code, s = call(base, "/api/sheets", {"file_path": xl})
    assert code == 200, s
    f = {x["nom"]: x for x in s["feuilles"]}
    assert f["Données"]["recommandee"] and f["Données"]["lignes"] == 120 and f["Données"]["colonnes"] == 4
    assert not f["Sommaire"]["exploitable"] and not f["Feuil3"]["exploitable"]
    assert f["Districts"]["masquee"] and f["Districts"]["exploitable"]
    # Sans choix explicite : la feuille recommandée, en-tête trouvé sous le titre
    code, a = call(base, "/api/analyze", {"file_path": xl})
    assert code == 200 and a["rows"] == 120 and a["columns"] == 4, a
    assert a["feuille"]["lue"] == "Données" and "Sommaire" in a["feuille"]["autres"]
    # Extraction : une feuille = un fichier de travail, utilisable par toutes les étapes
    code, x = call(base, "/api/sheets/extract", {"file_path": xl, "sheets": ["Données", "Districts", "Inconnue"]})
    assert code == 200 and [v["feuille"] for v in x["fichiers"]] == ["Données", "Districts"], x
    code, a = call(base, "/api/analyze", {"file_path": x["fichiers"][1]["path"]})
    assert a["rows"] == 6 and "feuille" not in a
    monkeypatch.setattr(server, "_llm_config", lambda: None)
    rel = server.architect_propose([v["path"] for v in x["fichiers"]])
    assert any("district" in str(r).lower() for r in rel["proposition"].get("relations") or []), rel
    for v in x["fichiers"]:
        os.remove(v["path"])
    assert f["Janv"]["groupe"] == f["Fevr"]["groupe"] == "Janv" and "groupe" not in f["Données"]
    code, x = call(base, "/api/sheets/extract", {"file_path": xl, "sheets": ["Janv", "Fevr"], "empiler": True})
    assert code == 200 and x["fichiers"][0]["lignes"] == 10, x
    code, a = call(base, "/api/analyze", {"file_path": x["fichiers"][0]["path"]})
    assert a["rows"] == 10 and a["columns_detail"][0]["name"] == "feuille", a
    os.remove(x["fichiers"][0]["path"])
    assert call(base, "/api/sheets", {"file_path": str(tmp_path / "absent.xlsx")})[1].get("error")


def test_architecte_cli_toutes_les_feuilles(tmp_path, capsys):
    import data_architect as da
    xl = _classeur(tmp_path / "rapport.xlsx")
    assert da.best_sheet(xl) == "Données"
    df, info = da.load_table(xl)
    assert info["feuille_lue"] == "Données" and len(df) == 120 and list(df.columns)[0] == "district"
    assert da.main(["profile", xl, "--all-sheets", "--out", str(tmp_path / "r.json")]) == 0
    import json
    rep = json.loads((tmp_path / "r.json").read_text(encoding="utf-8"))
    assert set(rep["fichiers"]) == {"rapport_Données", "rapport_Districts", "rapport_Janv", "rapport_Fevr"}
    assert da.main(["profile", xl + "#Districts", "--out", str(tmp_path / "d.json")]) == 0
